from pathlib import Path

from docx import Document


def set_para_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def cell_text(cell):
    return " ".join(cell.text.split())


def replace_cell_text(cell, text):
    if cell.paragraphs:
        set_para_text(cell.paragraphs[0], text)
        for paragraph in cell.paragraphs[1:]:
            set_para_text(paragraph, "")
    else:
        cell.text = text


path = Path("E:/qq下载路径/第八组项目详细设计 (1).docx")
doc = Document(str(path))

for paragraph in doc.paragraphs:
    text = "".join(run.text for run in paragraph.runs)
    if "OptimisticLockerInnerInterceptor" in text:
        set_para_text(
            paragraph,
            "当前V1.0版本已在报销单主表 fk_reim_main 中加入 version 乐观锁字段，"
            "并通过 MyBatis-Plus @Version 与 OptimisticLockerInnerInterceptor 防止并发覆盖。"
            "前端打开报销单详情时保存当前 version，保存草稿时随 ReimbursementDraftSaveDTO 一并提交；"
            "后端保存前比对数据库最新 version，更新影响行数为 0 或版本不一致时返回 "
            "BusinessException(409, \"报销单已被他人修改，请刷新后重试\")。"
            "报销单号仍使用数据库自增ID生成，避免并发冲突。",
        )

for table in doc.tables:
    for row in table.rows:
        cells = row.cells
        texts = [cell_text(cell) for cell in cells]
        joined = " | ".join(texts)
        if len(cells) >= 3 and texts[0] == "MyBatis-Plus":
            replace_cell_text(cells[2], "提供Lambda查询构造器、分页插件和乐观锁插件的ORM框架")
        if len(cells) >= 4 and texts[2] == "MybatisPlusConfig":
            replace_cell_text(cells[3], "分页插件与乐观锁插件配置")
        if len(cells) >= 4 and texts[0] == "fk_reim_main":
            replace_cell_text(cells[2], "reim_no, version, 状态/金额/人员/部门/公司/业务类型")
        if "ReimbursementDraftSaveDTO" in joined:
            target_index = next((idx for idx, text in enumerate(texts) if "ReimbursementDraftSaveDTO" in text), 1)
            replace_cell_text(
                cells[target_index],
                "ReimbursementDraftSaveDTO:{reimbursementTitle, reimburserId, reimDepartmentId, "
                "reimCompanyId, businessTypeId, businessTripReason, remarks, version, "
                "trips:[{travelerId, departCityId, arriveCityId, departDate, arriveDate, "
                "tripDescription, subsidyDays:[{subsidyDate, mealSelected, mealAmount, ...}]}], "
                "allocations:[{companyId, projectId, allocationRatio, allocationAmount}]}",
            )

doc.save(str(path))
