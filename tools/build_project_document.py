from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


WORKSPACE = Path(r"E:\codelocation\IDEA\travel")
OUT_DIR = WORKSPACE / "outputs" / "project-document"
OUT_FILE = OUT_DIR / "胜意科技项目文档.docx"

CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(32, 48, 63)
GRAY = RGBColor(86, 96, 108)
LIGHT_FILL = "E8EEF5"
MUTED_FILL = "F2F4F7"
CALL_OUT_FILL = "F4F6F9"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False, size: float = 9.0, color: RGBColor | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run("" if text is None else str(text))
    set_run(run, size=size, bold=bold, color=color or INK)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run(run, font: str = "Calibri", size: float | None = None, bold: bool | None = None,
            color: RGBColor | None = None) -> None:
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_cell_margins(cell, top: int = 80, bottom: int = 80, start: int = 120, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: Sequence[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    grid = tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths[index]))
            set_cell_margins(cell)


def add_table(doc: Document, headers: Sequence[str], rows: Iterable[Sequence[str]], widths: Sequence[int],
              font_size: float = 8.6, header_fill: str = LIGHT_FILL) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_text(cell, header, bold=True, size=font_size, color=DARK_BLUE)
        set_cell_shading(cell, header_fill)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            set_cell_text(cell, value, size=font_size)
    set_table_geometry(table, widths)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)


def add_kv_table(doc: Document, rows: Sequence[tuple[str, str]]) -> None:
    add_table(doc, ["项目", "说明"], rows, [2200, 7160], font_size=9.0)


def paragraph(doc: Document, text: str, bold_prefix: str | None = None, style: str | None = None) -> None:
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run(first, size=11, bold=True, color=INK)
        rest = p.add_run(text[len(bold_prefix):])
        set_run(rest, size=11, color=INK)
    else:
        run = p.add_run(text)
        set_run(run, size=11, color=INK)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25


def bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_run(run, size=10.5, color=INK)


def number_item(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_run(run, size=10.5, color=INK)


def heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)


def callout(doc: Document, title: str, body: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, CALL_OUT_FILL)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(title)
    set_run(run, size=10.2, bold=True, color=DARK_BLUE)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.15
    run = p2.add_run(body)
    set_run(run, size=9.6, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.page_height = Inches(11)
    section.page_width = Inches(8.5)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15

    header = section.header
    hp = header.paragraphs[0]
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run("胜意科技项目文档 | 差旅费用报销")
    set_run(run, size=9, color=GRAY)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("内部设计文档 | 当前范围：列表、录入、草稿、提交、撤回、作废")
    set_run(run, size=8.5, color=GRAY)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(34)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run("胜意科技项目文档")
    set_run(run, size=26, bold=True, color=INK)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(22)
    run = p.add_run("差旅费用报销后端详细设计与接口定义")
    set_run(run, size=15, color=DARK_BLUE)
    add_table(doc, ["文档属性", "内容"], [
        ["所属项目", "差旅费用报销单后端"],
        ["文档类型", "详细设计说明 + API 接口定义 + 数据库设计"],
        ["文档状态", "草稿 / 评审稿"],
        ["版本", "V1.0"],
        ["编制日期", "2026-05-22"],
        ["编制依据", "差旅报销单概要设计、详细设计规范模板、API 接口文档模板、当前后端实现"],
        ["实现边界", "仅实现报销单列表、报销单录入、保存草稿、关闭自动保留、提交、撤回、编辑、本人作废；不含审批、财务复核、结算和出纳流程"],
    ], [2200, 7160], font_size=9.2)
    callout(doc, "范围冻结说明",
            "状态机采用需求确认后的三态：草稿(code=0)、已完成(code=1)、已作废(code=2)。基础数据由数据库表维护；前端暂未生成，后端通过草稿保存接口支撑关闭时自动保留。")
    heading(doc, "修订记录", 1)
    add_table(doc, ["日期", "版本", "修改章节", "修改说明", "作者"], [
        ["2026-05-22", "V1.0", "全量", "依据需求确认和当前后端实现形成项目详细设计与接口定义文档", "Codex"],
    ], [1500, 900, 1800, 4160, 1000], font_size=8.8)
    heading(doc, "目录", 1)
    for item in [
        "1 开发须知与范围说明",
        "2 术语定义",
        "3 功能性需求描述",
        "4 功能详细设计",
        "5 技术实现设计",
        "6 关键技术点",
        "7 数据库设计",
        "8 WBS 与验证",
        "附录 A API 接口定义明细",
        "附录 B 数据字典与配置说明",
    ]:
        bullet(doc, item)
    doc.add_page_break()


def add_scope_and_terms(doc: Document) -> None:
    heading(doc, "1 开发须知与范围说明", 1)
    paragraph(doc, "本文档用于指导胜意科技差旅费用报销单后端开发、接口联调、数据库初始化和后续前端接入。文档结构参考《详细设计规范模板》，接口字段表参考《API接口接口文档模版》。")
    heading(doc, "1.1 本期目标", 2)
    for text in [
        "提供基于 JWT 的后端访问控制和统一响应格式。",
        "提供员工、部门、公司、业务类型、城市、项目等基础数据查询接口。",
        "提供报销单列表查询、详情查询、草稿创建、草稿更新、提交、撤回、作废接口。",
        "支持补录行程、补助日历、费用合计、费用归属及分摊的聚合保存。",
        "输出 MySQL 建表 SQL、测试数据、索引和外键关系，并保留后续前端联调所需 DTO/VO 契约。",
    ]:
        bullet(doc, text)
    heading(doc, "1.2 本期不做", 2)
    for text in [
        "不实现审批人、财务初审、财务共享复核、财务共享结算、出纳支付等流程图后续环节。",
        "不实现 BPM、发票识别、预算冻结、冲抵借款、结算单、会计凭证和第三方推送。",
        "不生成前端页面；前端自动保存策略由后续页面在关闭或离开前调用草稿保存接口完成。",
        "不引入 MQ、Redis、Job 的运行依赖；只给出适用性分析和后续扩展建议。",
    ]:
        bullet(doc, text)
    heading(doc, "1.3 非功能性要求", 2)
    add_table(doc, ["维度", "设计要求", "当前落点"], [
        ["可靠性", "草稿保存、提交、撤回、作废均要保持主子表一致；异常统一返回。", "报销单聚合保存与状态流转使用本地事务；统一异常处理。"],
        ["性能", "列表分页、查询条件索引、基础数据轻量返回。", "MyBatis-Plus 分页；主表状态、时间、报销人、公司、所有人索引。"],
        ["可维护性", "DTO/VO/Entity 分层，基础数据与报销域分包。", "auth、masterdata、reimbursement、common、config 分模块。"],
        ["安全性", "认证后访问；本人仅操作本人单据；敏感配置不应硬编码生产值。", "JWT Filter、owner_user_id 权限校验、环境变量配置占位符。"],
    ], [1500, 3660, 4200], font_size=8.4)
    heading(doc, "2 术语定义", 1)
    add_table(doc, ["术语", "定义"], [
        ["报销单主单", "承载报销标题、报销人、部门、费用归属公司、业务类型、出差事由、汇总金额、备注和状态的聚合根。"],
        ["补录行程", "员工手工补录的出行人、出发城市、到达城市、出发到达日期和行程说明。"],
        ["补助日历", "按行程日期逐天生成的餐补、交通补助、通讯补助勾选与金额明细。"],
        ["费用分摊", "将补助总金额按费用归属公司、项目、比例和金额拆分。"],
        ["草稿", "未完成或未提交的可编辑状态，code=0。"],
        ["已完成", "提交成功后的状态，code=1；当前范围内可撤回重新编辑。"],
        ["已作废", "由单据所有人作废后的终止状态，code=2，不允许继续编辑。"],
        ["基础数据", "数据库维护的员工、部门、公司、业务类型、城市和项目数据。"],
    ], [2200, 7160], font_size=8.8)


def add_requirements_and_functions(doc: Document) -> None:
    heading(doc, "3 功能性需求描述", 1)
    heading(doc, "3.1 业务背景", 2)
    paragraph(doc, "员工因公出差后需要报账结算。报销单录入过程既包含基础信息，又包含由行程推导的补助日历、费用合计和分摊信息。人工处理时容易出现日期重复、金额超过标准、分摊不平和提交信息不完整等问题，因此后端需要把这些校验沉淀为统一业务规则。")
    heading(doc, "3.2 用户角色", 2)
    add_table(doc, ["角色", "本期权限", "备注"], [
        ["报销提单人 / 员工", "登录、查询本人单据、保存草稿、编辑、提交、撤回、本人作废。", "本期实际实现角色。"],
        ["管理员", "可扩展全量查看权限。", "当前代码保留 ADMIN 判断，基础数据维护接口未实现。"],
        ["审批与财务角色", "本期不实现。", "需求流程图存在，但功能清单和确认范围已排除。"],
    ], [2200, 4200, 2960], font_size=8.6)
    heading(doc, "3.3 业务流程", 2)
    for text in [
        "用户登录后获取 JWT，前端在后续请求中携带 Bearer Token。",
        "用户进入报销单列表，按单号、标题、事由、费用归属公司、部门、报销人、业务类型分页查询本人单据。",
        "用户新增报销单，填写基础信息并补录行程；行程保存时后端生成补助汇总与逐日补助日历。",
        "用户勾选补助项、填写费用分摊；点击保存草稿或关闭页面自动保存时，后端保存整页聚合快照。",
        "用户提交时，后端执行主单必填、行程非空、行程日期重叠、分摊比例和分摊金额校验。",
        "提交成功后状态变为已完成；用户可以撤回回到草稿，也可以按本人权限作废。",
    ]:
        number_item(doc, text)
    heading(doc, "3.4 页面需求与后端映射", 2)
    add_table(doc, ["页面 / 弹窗", "关键控件", "后端映射"], [
        ["报销单列表页", "查询区、列表列、分页、新增入口", "GET /api/reimbursements；GET /api/master/*"],
        ["报销单录入页", "基础信息、行程、补助、合计、分摊、备注、固定底部按钮", "POST /api/reimbursements/drafts；PUT /api/reimbursements/{id}/draft"],
        ["补录行程弹窗", "出行人、城市、日期、行程说明", "聚合保存在 trips 数组；后端生成 subsidy 和 subsidyDays"],
        ["补助日历弹窗", "餐补、交通、通讯勾选和金额", "subsidyDays 明细数组；金额受标准上限约束"],
        ["状态操作", "提交、撤回、作废", "POST /submit、POST /withdraw、POST /void"],
    ], [1900, 3760, 3700], font_size=8.2)
    heading(doc, "3.5 功能模块清单", 2)
    add_table(doc, ["模块", "一级功能", "二级功能", "实现说明"], [
        ["认证与安全", "登录", "JWT 发放", "用户名密码校验后返回 Bearer Token。"],
        ["基础数据", "下拉控件数据", "员工、部门、公司、业务类型树、城市、项目", "数据库读取，业务表保留名称快照。"],
        ["报销单列表", "分页查询", "多条件过滤", "仅本人数据；管理员扩展可全量。"],
        ["报销单录入", "草稿保存", "整页聚合保存", "关闭自动保留复用草稿接口。"],
        ["行程补助", "补录行程", "逐日补助日历", "到达城市决定城市等级与餐补标准。"],
        ["费用处理", "合计与分摊", "比例、金额校验", "提交时比例=1，分摊金额=补助总金额。"],
        ["状态流转", "提交、撤回、作废", "三态控制", "作废后不可继续编辑。"],
    ], [1500, 1700, 2600, 3560], font_size=8.2)


def add_function_design(doc: Document) -> None:
    heading(doc, "4 功能详细设计", 1)
    function_sections = [
        ("4.1 登录与权限",
         ["功能内容：提供用户名密码登录，登录成功后生成 JWT；除登录接口外所有接口均需认证。",
          "实现逻辑：AuthService 查询 sys_user，使用 PBKDF2 校验 password_hash；JwtTokenService 使用 HMAC-SHA256 签发包含 uid、eid、roles、iat、exp 的令牌；JwtAuthenticationFilter 将令牌解析为 CurrentUser。",
          "异常处置：用户名密码错误返回业务失败；令牌格式错误、签名错误、过期令牌统一拒绝。"]),
        ("4.2 基础数据查询",
         ["功能内容：为页面下拉框提供公司、部门、员工、业务类型树、城市、项目。",
          "实现逻辑：MasterDataService 读取基础数据表；业务类型按 parent_id 组装树；城市返回 cityType 用于前端展示与后端补助计算。",
          "异常处置：查询接口本身只读；不存在的数据在报销单保存时由领域服务再次校验。"]),
        ("4.3 报销单列表",
         ["功能内容：支持分页、条件查询、展示状态和补助金额。",
          "实现逻辑：按 current、size 分页；支持 reimNo、title、reason、companyId、departmentId、reimburserId、businessTypeId 条件；非管理员追加 owner_user_id 条件。",
          "异常处置：分页大小做上下界限制；无数据返回空 records，不返回异常。"]),
        ("4.4 草稿创建与关闭自动保留",
         ["功能内容：新增草稿、更新草稿、支持页面关闭前自动保存未提交数据。",
          "实现逻辑：草稿允许主单字段不完整；首次保存插入主表后生成 SY + 10位主键流水号；更新草稿时采用聚合快照替换子表数据。",
          "异常处置：已作废单据拒绝编辑；非本人单据拒绝操作；草稿中已提供的行程仍需满足日期、城市、说明和补助金额边界。"]),
        ("4.5 补录行程与补助日历",
         ["功能内容：记录手工行程，并据此生成补助汇总和逐日补助明细。",
          "实现逻辑：行程要求出行人、出发城市、到达城市、日期、说明完整；到达城市 cityType 决定餐补标准；每个日期生成一条 subsidy_day，交通与通讯标准固定为40。",
          "异常处置：到达日期早于出发日期、到达日期晚于当前日期、同一人日期区间重叠、补助日期不在行程范围、同一补助日期重复均拒绝。"]),
        ("4.6 费用合计与费用分摊",
         ["功能内容：按补助日历回算餐补、交通、通讯和总额；将总额分摊到公司与项目。",
          "实现逻辑：Subsidy 汇总每日日历金额；Reimbursement 主表持久化 subsidy_total、meal_allowance、transportation_allowance、phone_allowance；Allocation 记录比例和值。",
          "异常处置：金额不可为负；单项补助金额不可大于标准；提交时校验分摊比例合计100%，金额合计与主单补助总额相等。"]),
        ("4.7 提交、撤回与作废",
         ["功能内容：实现当前确认的三态状态流转。",
          "实现逻辑：草稿或已完成单据可编辑；submit 完成严格校验并置为已完成；withdraw 仅对已完成单据置回草稿；void 由所有人将未作废单据置为已作废。",
          "异常处置：重复作废、非已完成撤回、作废后编辑均返回业务异常。"]),
    ]
    for title, lines in function_sections:
        heading(doc, title, 2)
        for line in lines:
            paragraph(doc, line, bold_prefix=line.split("：")[0] + "：")
    heading(doc, "4.8 状态机", 2)
    add_table(doc, ["当前状态", "动作", "目标状态", "校验"], [
        ["草稿(0)", "保存草稿", "草稿(0)", "可部分保存；已提供的行程和金额必须合法。"],
        ["草稿(0)", "提交", "已完成(1)", "主单必填、行程非空、分摊比例与金额合计通过。"],
        ["已完成(1)", "编辑后保存", "已完成(1)", "当前实现允许编辑已完成单据，前端可按业务需要先撤回再编辑。"],
        ["已完成(1)", "撤回", "草稿(0)", "仅本人、仅已完成可撤回。"],
        ["草稿(0) / 已完成(1)", "作废", "已作废(2)", "仅本人可作废。"],
        ["已作废(2)", "编辑 / 提交", "拒绝", "已作废报销单不可编辑。"],
    ], [1600, 1800, 1600, 4360], font_size=8.4)
    heading(doc, "4.9 业务校验矩阵", 2)
    add_table(doc, ["校验点", "草稿保存", "提交", "规则"], [
        ["主单基础字段", "可为空", "必填", "标题、报销人、部门、公司、业务类型、出差事由。"],
        ["文本长度", "校验", "校验", "标题与事由<=500，备注<=1000，行程说明<=500。"],
        ["行程字段", "已传行程必填", "已传行程必填且至少一条", "出行人、城市、日期、说明。"],
        ["行程日期", "校验", "校验", "到达日期>=出发日期，到达日期<=当前日期，同一人区间不可重叠。"],
        ["补助金额", "校验", "校验", "仅选中项计入；输入金额在0到标准金额之间。"],
        ["分摊", "可空", "必填且平衡", "至少一行；比例合计=1；金额合计=补助总金额。"],
    ], [2000, 1400, 1600, 4360], font_size=8.3)


def add_technical_design(doc: Document) -> None:
    heading(doc, "5 技术实现设计", 1)
    heading(doc, "5.1 系统结构设计", 2)
    add_table(doc, ["层级", "职责", "当前包路径 / 组件"], [
        ["接入层", "REST 接口、参数绑定、认证主体传递", "controller、JwtAuthenticationFilter"],
        ["应用服务层", "业务编排、事务边界、聚合保存", "AuthService、MasterDataService、ReimbursementService"],
        ["数据访问层", "Entity 与 MyBatis-Plus BaseMapper", "entity、mapper"],
        ["公共能力层", "统一返回、异常、分页、状态枚举", "common.result、common.exception、reimbursement.enums"],
        ["配置层", "Security、JWT 配置、MyBatis-Plus 分页", "config.security、config.mybatis"],
    ], [1700, 3100, 4560], font_size=8.5)
    paragraph(doc, "系统采用前后端分离后端服务形态。请求先经过 Spring Security 与 JWT 过滤器，Controller 只负责 REST 契约，核心业务规则集中在 Service；数据库访问统一走 MyBatis-Plus Mapper。")
    heading(doc, "5.2 技术栈", 2)
    add_kv_table(doc, [
        ["Java", "17 语法级别，当前本机测试运行于 Java 21。"],
        ["Web 框架", "Spring Boot 4.0.6 + Spring Web MVC。"],
        ["ORM", "MyBatis-Plus 3.5.15，分页插件另引入 mybatis-plus-jsqlparser。"],
        ["数据库", "MySQL 8 目标库；测试使用 H2 MySQL MODE。"],
        ["认证", "Spring Security + 自定义 JWT HMAC-SHA256。"],
        ["校验", "Jakarta Validation + 领域服务业务校验。"],
        ["工程辅助", "Lombok、Maven、SQL 初始化脚本。"],
    ])
    heading(doc, "5.3 核心类设计", 2)
    add_table(doc, ["类", "职责", "关键方法 / 备注"], [
        ["ReimbursementController", "报销接口入口", "page、detail、createDraft、saveDraft、submit、withdraw、voidByOwner"],
        ["ReimbursementService", "报销聚合根应用服务", "聚合保存、补助计算、分摊校验、状态流转、权限判断"],
        ["MasterDataService", "下拉基础数据查询", "companies、departments、employees、businessTypeTree、cities、projects"],
        ["AuthService", "登录服务", "查询用户、校验密码、签发 JWT"],
        ["JwtTokenService", "JWT 编解码与签名", "createToken、parseToken"],
        ["GlobalExceptionHandler", "异常转 Result", "参数异常、认证异常、业务异常、兜底异常"],
    ], [2500, 3100, 3760], font_size=8.2)
    heading(doc, "5.4 与前端交互", 2)
    add_table(doc, ["前端动作", "调用接口", "返回数据关注点"], [
        ["登录", "POST /api/auth/login", "token、tokenType、expiresInMinutes"],
        ["初始化下拉", "GET /api/master/*", "基础数据 id、编号、名称、业务类型树、城市等级"],
        ["列表查询", "GET /api/reimbursements", "PageResult records、total、current、size"],
        ["新增并保存草稿", "POST /api/reimbursements/drafts", "返回详情、生成 reimNo"],
        ["关闭自动保留 / 手工保存", "PUT /api/reimbursements/{id}/draft", "页面整页数据快照"],
        ["提交 / 撤回 / 作废", "POST /submit、/withdraw、/void", "状态 code 与状态名"],
    ], [2400, 3000, 3960], font_size=8.3)
    heading(doc, "5.5 与第三方交互", 2)
    paragraph(doc, "本期功能没有第三方系统调用。流程图中潜在 BPM、发票、结算、凭证和消息推送能力保留为后续扩展点，当前接口不会对外部系统产生副作用。")
    heading(doc, "5.6 项目目录结构", 2)
    add_table(doc, ["目录", "说明"], [
        ["auth", "登录、密码校验、用户实体、登录 DTO/VO。"],
        ["common", "Result、PageResult、业务异常和全局异常处理。"],
        ["config.security", "JWT 配置、过滤器、SecurityFilterChain、CurrentUser。"],
        ["config.mybatis", "MyBatis-Plus 分页拦截器。"],
        ["masterdata", "员工、部门、公司、业务类型、城市、项目查询。"],
        ["reimbursement", "报销 DTO、VO、Entity、Mapper、Service、Controller。"],
        ["resources/db", "schema.sql 建表脚本，data.sql 测试数据。"],
        ["test", "上下文启动和报销核心流测试。"],
    ], [2800, 6560], font_size=8.6)


def add_key_technical_points(doc: Document) -> None:
    heading(doc, "6 关键技术点", 1)
    add_table(doc, ["主题", "分析结论", "本期设计"], [
        ["并发编程", "当前流程无异步批处理，核心风险是重复保存和状态更新。", "单次聚合事务处理；后续可引入乐观锁 version 字段控制并发覆盖。"],
        ["事务控制", "主单与行程、补助、日历、分摊必须一致。", "createDraft、saveDraft、submit、withdraw、void 使用 @Transactional。"],
        ["Job", "本期无超时任务和补偿扫描。", "不引入；后续可用于草稿清理、失败推送重试、提醒。"],
        ["Redis", "基础数据和 JWT 黑名单有缓存价值。", "本期不引入；后续缓存 key 要有前缀、TTL 和刷新策略。"],
        ["MQ", "本期无外部异步事件。", "不引入；后续对 BPM、消息通知、结算推送可发领域事件并做幂等。"],
        ["权限控制", "本人单据隔离是当前最重要的数据权限。", "owner_user_id 查询过滤、操作校验；ADMIN 留扩展口。"],
        ["敏感信息", "数据库密码和 JWT secret 不宜写生产明文。", "application.properties 使用环境变量占位符和开发默认值。"],
        ["错误码", "需要统一失败结构便于前端提示。", "Result.code + message；当前业务异常主要用 400/401/403/404/500。"],
        ["异动日志", "提交、撤回、作废属于关键状态异动。", "当前未单独建审计表；建议后续补 operation_log。"],
        ["大数据量", "列表和明细随时间增长。", "主表索引、分页；后续可做归档和按状态/时间治理。"],
        ["重试机制", "本期数据库内事务不需要异步重试。", "未来第三方接口调用应设计幂等键和补偿重试。"],
    ], [1500, 3600, 4260], font_size=7.9)
    doc.add_page_break()
    heading(doc, "6.1 配置说明", 2)
    add_table(doc, ["配置项", "环境变量", "默认值", "说明"], [
        ["spring.datasource.url", "TRAVEL_DB_URL", "jdbc:mysql://localhost:3306/travel?...", "数据库连接。"],
        ["spring.datasource.username", "TRAVEL_DB_USERNAME", "root", "数据库账号。"],
        ["spring.datasource.password", "TRAVEL_DB_PASSWORD", "root", "数据库密码，生产必须覆盖。"],
        ["travel.jwt.secret", "TRAVEL_JWT_SECRET", "change-this-secret-before-production-change-this-secret", "JWT HMAC 密钥，生产必须覆盖。"],
        ["travel.jwt.expiration-minutes", "TRAVEL_JWT_EXPIRATION_MINUTES", "480", "令牌有效期分钟数。"],
    ], [2300, 2200, 2860, 2000], font_size=8.1)
    callout(doc, "安全提醒", "默认 JWT 密钥与默认数据库口令只适用于本地开发占位。生产部署必须通过环境变量或安全配置中心覆盖。")


DB_TABLES = [
    ("sys_company", "费用归属公司基础表", [
        ("id", "BIGINT", "PK", "自增主键"),
        ("company_no", "VARCHAR(32)", "UK/NOT NULL", "公司编号"),
        ("company_name", "VARCHAR(128)", "NOT NULL", "公司名称"),
        ("created_at", "DATETIME", "NOT NULL", "创建时间"),
    ]),
    ("sys_department", "报销部门基础表", [
        ("id", "BIGINT", "PK", "自增主键"),
        ("department_no", "VARCHAR(32)", "UK/NOT NULL", "部门编号"),
        ("department_name", "VARCHAR(128)", "NOT NULL", "部门名称"),
        ("created_at", "DATETIME", "NOT NULL", "创建时间"),
    ]),
    ("sys_employee", "员工基础表", [
        ("id", "BIGINT", "PK", "自增主键"),
        ("employee_no", "VARCHAR(32)", "UK/NOT NULL", "员工工号"),
        ("employee_name", "VARCHAR(64)", "NOT NULL", "员工姓名"),
        ("department_id", "BIGINT", "FK", "所属部门"),
        ("company_id", "BIGINT", "FK", "所属公司"),
    ]),
    ("sys_user", "登录用户表", [
        ("id", "BIGINT", "PK", "自增主键"),
        ("username", "VARCHAR(64)", "UK/NOT NULL", "登录名"),
        ("password_hash", "VARCHAR(255)", "NOT NULL", "PBKDF2 密码摘要"),
        ("employee_id", "BIGINT", "FK/NOT NULL", "绑定员工"),
        ("roles", "VARCHAR(255)", "NOT NULL", "角色集合"),
        ("enabled", "TINYINT", "NOT NULL", "启用标志"),
    ]),
    ("biz_business_type", "业务类型树表", [
        ("id", "BIGINT", "PK", "自增主键"),
        ("business_type_no", "VARCHAR(32)", "UK/NOT NULL", "业务类型编号"),
        ("business_type_name", "VARCHAR(128)", "NOT NULL", "业务类型名称"),
        ("parent_id", "BIGINT", "FK", "父级类型"),
        ("leaf_flag", "TINYINT", "NOT NULL", "叶子节点标志"),
    ]),
    ("biz_city", "城市基础表", [
        ("id", "BIGINT", "PK", "自增主键"),
        ("city_no", "VARCHAR(32)", "UK/NOT NULL", "城市编号"),
        ("city_name", "VARCHAR(64)", "NOT NULL", "城市名称"),
        ("city_type", "TINYINT", "NOT NULL", "1一线、2二线、3三线"),
    ]),
    ("biz_project", "项目基础表", [
        ("id", "BIGINT", "PK", "自增主键"),
        ("project_no", "VARCHAR(64)", "UK/NOT NULL", "项目编号"),
        ("project_name", "VARCHAR(128)", "NOT NULL", "项目名称"),
    ]),
    ("fk_reim_main", "报销单主表", [
        ("id", "BIGINT", "PK", "自增主键"),
        ("reim_no", "VARCHAR(32)", "UK", "SY 前缀流水号"),
        ("reimbursement_title", "VARCHAR(500)", "NULL", "报销标题"),
        ("reimburser_id/no/name", "BIGINT/VARCHAR", "快照", "报销人"),
        ("reim_department_id/no/name", "BIGINT/VARCHAR", "快照", "报销部门"),
        ("reim_company_id/no/name", "BIGINT/VARCHAR", "快照", "费用归属公司"),
        ("business_type_id/no/name", "BIGINT/VARCHAR", "快照", "业务类型"),
        ("business_trip_reason", "VARCHAR(500)", "NULL", "出差事由"),
        ("subsidy_total", "DECIMAL(18,2)", "NOT NULL", "补助总金额"),
        ("meal_allowance", "DECIMAL(18,2)", "NOT NULL", "餐费补助汇总"),
        ("transportation_allowance", "DECIMAL(18,2)", "NOT NULL", "交通补助汇总"),
        ("phone_allowance", "DECIMAL(18,2)", "NOT NULL", "通讯补助汇总"),
        ("remarks", "VARCHAR(1000)", "NULL", "备注"),
        ("status", "TINYINT", "NOT NULL", "0草稿、1已完成、2已作废"),
        ("owner_user_id", "BIGINT", "FK/NOT NULL", "单据所有人"),
    ]),
    ("fk_reim_trip", "补录行程表", [
        ("id", "BIGINT", "PK", "自增主键"),
        ("main_id", "BIGINT", "FK/NOT NULL", "所属主单"),
        ("traveler_id/no/name", "BIGINT/VARCHAR", "NOT NULL", "出行人快照"),
        ("depart_city_id/no/name", "BIGINT/VARCHAR", "NOT NULL", "出发城市快照"),
        ("arrive_city_id/no/name", "BIGINT/VARCHAR", "NOT NULL", "到达城市快照"),
        ("depart_date", "DATE", "NOT NULL", "出发日期"),
        ("arrive_date", "DATE", "NOT NULL", "到达日期"),
        ("trip_description", "VARCHAR(500)", "NOT NULL", "行程说明"),
    ]),
    ("fk_reim_subsidy", "补助汇总表", [
        ("id", "BIGINT", "PK", "自增主键"),
        ("main_id", "BIGINT", "FK", "所属主单"),
        ("trip_id", "BIGINT", "UK/FK", "一条行程一条补助汇总"),
        ("subsidy_days", "INT", "NOT NULL", "补助天数"),
        ("apply_amount", "DECIMAL(18,2)", "NOT NULL", "选中标准总额"),
        ("subsidy_amount", "DECIMAL(18,2)", "NOT NULL", "实际补助金额"),
        ("meal_amount/transportation_amount/phone_amount", "DECIMAL(18,2)", "NOT NULL", "分项汇总"),
    ]),
    ("fk_reim_subsidy_day", "补助日历明细表", [
        ("id", "BIGINT", "PK", "自增主键"),
        ("subsidy_id", "BIGINT", "FK", "所属补助汇总"),
        ("subsidy_date", "DATE", "UK PART", "补助日期"),
        ("weekday_name", "VARCHAR(16)", "NOT NULL", "星期显示"),
        ("city_id/city_name", "BIGINT/VARCHAR", "NOT NULL", "补助城市"),
        ("meal_standard_amount", "DECIMAL(18,2)", "NOT NULL", "餐补标准"),
        ("transportation_standard_amount", "DECIMAL(18,2)", "NOT NULL", "交通标准"),
        ("phone_standard_amount", "DECIMAL(18,2)", "NOT NULL", "通讯标准"),
        ("meal_selected/transportation_selected/phone_selected", "TINYINT", "NOT NULL", "勾选状态"),
        ("meal_amount/transportation_amount/phone_amount", "DECIMAL(18,2)", "NOT NULL", "实际金额"),
    ]),
    ("fk_reim_allocation", "费用分摊表", [
        ("id", "BIGINT", "PK", "自增主键"),
        ("main_id", "BIGINT", "FK", "所属主单"),
        ("company_id/no/name", "BIGINT/VARCHAR", "NOT NULL", "分摊公司快照"),
        ("project_id/no/name", "BIGINT/VARCHAR", "NULL", "项目快照"),
        ("allocation_ratio", "DECIMAL(8,6)", "NOT NULL", "0到1比例"),
        ("allocation_amount", "DECIMAL(18,2)", "NOT NULL", "分摊金额"),
        ("row_order", "INT", "NOT NULL", "行顺序"),
    ]),
]


def add_database_design(doc: Document) -> None:
    heading(doc, "7 数据库设计", 1)
    heading(doc, "7.1 ER 关系说明", 2)
    for text in [
        "fk_reim_main 是报销聚合根，一对多关联 fk_reim_trip 和 fk_reim_allocation。",
        "每条 fk_reim_trip 一对一生成 fk_reim_subsidy；每条 fk_reim_subsidy 一对多生成 fk_reim_subsidy_day。",
        "sys_user 绑定 sys_employee，报销主单 owner_user_id 控制本人权限。",
        "公司、部门、员工、业务类型、城市、项目是基础数据表；报销业务表同时保存 id、编号和名称快照，避免历史展示随基础数据更名而失真。",
    ]:
        bullet(doc, text)
    heading(doc, "7.2 表清单", 2)
    add_table(doc, ["表名", "中文名", "核心关系"], [
        ["sys_company", "公司基础表", "被员工、主单、分摊引用"],
        ["sys_department", "部门基础表", "被员工和主单引用"],
        ["sys_employee", "员工基础表", "被用户、主单、行程引用"],
        ["sys_user", "登录用户表", "拥有报销单"],
        ["biz_business_type", "业务类型树", "主单业务类型"],
        ["biz_city", "城市数据", "行程城市与补助城市"],
        ["biz_project", "项目数据", "分摊项目"],
        ["fk_reim_main", "报销主表", "聚合根"],
        ["fk_reim_trip", "补录行程", "主表子项"],
        ["fk_reim_subsidy", "补助汇总", "行程派生"],
        ["fk_reim_subsidy_day", "补助日历", "补助逐日明细"],
        ["fk_reim_allocation", "费用分摊", "主表子项"],
    ], [2600, 2600, 4160], font_size=8.4)
    heading(doc, "7.3 表定义", 2)
    for table_name, cn_name, fields in DB_TABLES:
        heading(doc, f"7.3 {table_name} - {cn_name}", 3)
        add_table(doc, ["字段", "类型", "约束", "说明"], fields, [2800, 2100, 1760, 2700], font_size=7.8)
    heading(doc, "7.4 索引设计", 2)
    add_table(doc, ["表", "索引", "目的"], [
        ["fk_reim_main", "uk_fk_reim_main_no", "保证报销单号唯一。"],
        ["fk_reim_main", "idx_fk_reim_main_status_time", "支持列表按状态和创建时间检索。"],
        ["fk_reim_main", "idx_fk_reim_main_reimburser / company / owner", "支持报销人、公司、本人数据权限过滤。"],
        ["fk_reim_trip", "idx_fk_reim_trip_main", "加速主单详情加载。"],
        ["fk_reim_trip", "idx_fk_reim_trip_traveler_dates", "辅助人员日期范围判断。"],
        ["fk_reim_subsidy", "uk_fk_reim_subsidy_trip", "保证一条行程一条补助汇总。"],
        ["fk_reim_subsidy_day", "uk_fk_reim_subsidy_day_date", "保证同一补助日期不重复。"],
        ["fk_reim_allocation", "idx_fk_reim_allocation_main", "加速主单分摊加载。"],
    ], [2200, 3400, 3760], font_size=8.1)
    heading(doc, "7.5 数据流向", 2)
    for text in [
        "基础数据 -> 页面下拉 -> 草稿请求 DTO -> 领域服务校验 -> 主表快照。",
        "补录行程 -> 到达城市等级 -> 补助日历标准金额 -> 日历勾选与实际金额 -> 补助汇总 -> 主表费用合计。",
        "主表费用合计 -> 分摊明细 -> 提交校验 -> 状态已完成。",
    ]:
        number_item(doc, text)


def add_wbs_and_verification(doc: Document) -> None:
    heading(doc, "8 WBS 与验证", 1)
    heading(doc, "8.1 开发任务拆解", 2)
    add_table(doc, ["阶段", "任务", "交付物", "状态"], [
        ["需求分析", "解读概要设计与模板，确认实现范围与歧义", "模块、ER、API、页面和权限分析", "已完成"],
        ["基础工程", "依赖、配置、统一 Result、统一异常、JWT", "common/config/auth 代码", "已完成"],
        ["数据库", "建表 SQL、测试数据、索引、外键", "schema.sql、data.sql", "已完成"],
        ["基础数据", "公司、部门、员工、业务类型、城市、项目接口", "masterdata 模块", "已完成"],
        ["报销域", "列表、草稿、提交、撤回、作废、补助计算、分摊校验", "reimbursement 模块", "已完成"],
        ["验证", "上下文启动与核心业务测试", "Maven test", "已完成"],
        ["后续", "前端页面、OpenAPI、审计日志与并发控制增强", "下一阶段任务", "待排期"],
    ], [1500, 3300, 3160, 1400], font_size=8.1)
    heading(doc, "8.2 已验证用例", 2)
    add_table(doc, ["用例", "验证点", "结果"], [
        ["TravelApplicationTests.contextLoads", "Spring Boot 测试上下文、H2 初始化、Mapper 与配置加载", "通过"],
        ["loginDraftSubmitWithdrawAndVoidFlowWorks", "登录、生成草稿单号、补助汇总、提交、撤回、作废", "通过"],
        ["overlappingTripsForSameTravelerAreRejected", "同一出行人日期重叠行程拦截", "通过"],
    ], [3100, 4560, 1700], font_size=8.4)
    callout(doc, "验证命令", "当前已执行 mvn.cmd test，测试总数 3，失败 0，错误 0。")


def api_intro(doc: Document, title: str, method_name: str, path: str, method: str, scene: str) -> None:
    heading(doc, title, 2)
    add_kv_table(doc, [
        ["接口调用地址", "ip + 端口 + " + path],
        ["方法名", method_name],
        ["HTTP 方法", method],
        ["接口路径", path],
        ["应用场景", scene],
        ["认证要求", "登录接口免认证，其余接口需 Authorization: Bearer {token}。"],
    ])


def api_params(doc: Document, rows: Sequence[Sequence[str]]) -> None:
    heading(doc, "接口入参", 3)
    add_table(doc, ["字段名", "类型", "是否必传", "备注", "格式 / 取值"], rows,
              [1900, 1800, 1300, 2860, 1500], font_size=7.7)


def api_output(doc: Document, rows: Sequence[Sequence[str]]) -> None:
    heading(doc, "接口出参", 3)
    add_table(doc, ["字段名", "类型", "是否返回", "备注", "格式 / 取值"], rows,
              [1900, 1800, 1300, 2860, 1500], font_size=7.7)


def add_api_appendix(doc: Document) -> None:
    heading(doc, "附录 A API 接口定义明细", 1)
    paragraph(doc, "接口统一返回 Result<T>，成功 code=0、message=success、data 为业务对象；异常由统一异常处理转换为 code 与 message。分页返回 PageResult<T>。")
    add_table(doc, ["公共响应字段", "类型", "说明"], [
        ["code", "Integer", "0 表示成功；400 参数或业务失败；401 身份失败；403 权限不足；404 资源不存在；500 系统异常。"],
        ["message", "String", "成功为 success，失败为可提示文本。"],
        ["data", "Object", "业务数据；无数据动作可为空。"],
    ], [2400, 1800, 5160], font_size=8.3)
    add_table(doc, ["分页字段", "类型", "说明"], [
        ["total", "Long", "总条数"],
        ["current", "Long", "当前页"],
        ["size", "Long", "每页大小"],
        ["records", "List<T>", "当前页记录"],
    ], [2400, 1800, 5160], font_size=8.3)

    api_intro(doc, "A.1 登录", "AUTH_login", "/api/auth/login", "POST", "用户输入用户名和密码，换取 JWT。")
    api_params(doc, [
        ["username", "String", "Y", "登录名", "非空"],
        ["password", "String", "Y", "密码", "非空"],
    ])
    api_output(doc, [
        ["token", "String", "Y", "JWT 值", "Bearer Token"],
        ["tokenType", "String", "Y", "令牌类型", "Bearer"],
        ["expiresInMinutes", "Long", "Y", "有效期分钟数", "默认480"],
    ])

    master_apis = [
        ("A.2 公司下拉", "MASTER_companies", "/api/master/companies", "公司下拉控件数据",
         [["id", "Long", "Y", "公司主键", ""], ["companyNo", "String", "Y", "公司编号", ""], ["companyName", "String", "Y", "公司名称", ""]]),
        ("A.3 部门下拉", "MASTER_departments", "/api/master/departments", "部门下拉控件数据",
         [["id", "Long", "Y", "部门主键", ""], ["departmentNo", "String", "Y", "部门编号", ""], ["departmentName", "String", "Y", "部门名称", ""]]),
        ("A.4 员工下拉", "MASTER_employees", "/api/master/employees", "员工控件数据",
         [["id", "Long", "Y", "员工主键", ""], ["employeeNo", "String", "Y", "员工工号", ""], ["employeeName", "String", "Y", "员工姓名", ""], ["departmentId", "Long", "Y", "部门主键", ""], ["companyId", "Long", "Y", "公司主键", ""]]),
        ("A.5 业务类型树", "MASTER_businessTypeTree", "/api/master/business-types/tree", "业务类型树型下拉数据",
         [["id", "Long", "Y", "业务类型主键", ""], ["businessTypeNo", "String", "Y", "业务类型编号", ""], ["businessTypeName", "String", "Y", "业务类型名称", ""], ["parentId", "Long", "Y", "父节点主键", "根节点为空"], ["leaf", "Boolean", "Y", "是否叶子", ""], ["children", "List<BusinessTypeTreeVO>", "Y", "子节点", "递归"]]),
        ("A.6 城市下拉", "MASTER_cities", "/api/master/cities", "城市控件和城市等级数据",
         [["id", "Long", "Y", "城市主键", ""], ["cityNo", "String", "Y", "城市编号", ""], ["cityName", "String", "Y", "城市名称", ""], ["cityType", "Integer", "Y", "城市等级", "1/2/3"]]),
        ("A.7 项目下拉", "MASTER_projects", "/api/master/projects", "费用分摊项目控件数据",
         [["id", "Long", "Y", "项目主键", ""], ["projectNo", "String", "Y", "项目编号", ""], ["projectName", "String", "Y", "项目名称", ""]]),
    ]
    for title, method_name, path, scene, output in master_apis:
        api_intro(doc, title, method_name, path, "GET", scene)
        api_params(doc, [["无", "-", "N", "无业务入参", "-"]])
        api_output(doc, output)

    api_intro(doc, "A.8 报销单分页查询", "REIM_page", "/api/reimbursements", "GET", "报销单列表页面按条件分页查询。")
    api_params(doc, [
        ["current", "Long", "N", "当前页", "默认1"],
        ["size", "Long", "N", "每页大小", "默认10，最大100"],
        ["reimNo", "String", "N", "报销单号模糊查询", "SY流水号"],
        ["title", "String", "N", "标题模糊查询", ""],
        ["reason", "String", "N", "事由模糊查询", ""],
        ["companyId", "Long", "N", "费用归属公司", ""],
        ["departmentId", "Long", "N", "报销部门", ""],
        ["reimburserId", "Long", "N", "报销人", ""],
        ["businessTypeId", "Long", "N", "业务类型", ""],
    ])
    api_output(doc, [
        ["records[].id", "Long", "Y", "报销单主键", ""],
        ["records[].reimNo", "String", "Y", "报销单号", "SY + 10位流水"],
        ["records[].status/statusName", "Integer/String", "Y", "状态", "0草稿/1已完成/2已作废"],
        ["records[].reimburserDisplay", "String", "Y", "报销人显示", "姓名[工号]"],
        ["records[].departmentDisplay", "String", "Y", "部门显示", "名称[编号]"],
        ["records[].reimCompanyName", "String", "Y", "费用归属公司", ""],
        ["records[].businessTypeName", "String", "Y", "业务类型", ""],
        ["records[].reimbursementTitle", "String", "Y", "报销标题", ""],
        ["records[].businessTripReason", "String", "Y", "出差事由", ""],
        ["records[].subsidyTotal", "Decimal", "Y", "补助总金额", "2位小数"],
        ["records[].creationTime", "DateTime", "Y", "创建时间", "yyyy-MM-dd HH:mm:ss"],
    ])

    api_intro(doc, "A.9 报销单详情", "REIM_detail", "/api/reimbursements/{id}", "GET", "编辑页或查看页加载整页聚合数据。")
    api_params(doc, [["id", "Long", "Y", "报销单主键", "Path"]])
    api_output(doc, [
        ["主单字段", "ReimbursementDetailVO", "Y", "单号、状态、基础信息、金额合计、备注和时间", "详见数据字典"],
        ["trips", "List<TripVO>", "Y", "行程集合", "内含 subsidy"],
        ["trips[].subsidy.days", "List<SubsidyDayVO>", "Y", "补助日历明细", "按日期排序"],
        ["allocations", "List<AllocationVO>", "Y", "费用分摊明细", "按 rowOrder 排序"],
    ])

    save_payload_rows = [
        ["reimbursementTitle", "String", "N", "报销标题", "<=500字"],
        ["reimburserId", "Long", "N", "报销人ID", "提交时必填"],
        ["reimDepartmentId", "Long", "N", "报销部门ID", "提交时必填"],
        ["reimCompanyId", "Long", "N", "费用归属公司ID", "提交时必填"],
        ["businessTypeId", "Long", "N", "业务类型ID", "提交时必填"],
        ["businessTripReason", "String", "N", "出差事由", "<=500字，提交时必填"],
        ["remarks", "String", "N", "备注", "<=1000字"],
        ["trips", "List<TripSaveDTO>", "N", "补录行程", "提交时至少1条"],
        ["trips[].travelerId", "Long", "Y*", "出行人", "行程存在时必填"],
        ["trips[].departCityId/arriveCityId", "Long", "Y*", "出发/到达城市", "行程存在时必填"],
        ["trips[].departDate/arriveDate", "LocalDate", "Y*", "出发/到达日期", "yyyy-MM-dd"],
        ["trips[].tripDescription", "String", "Y*", "行程说明", "<=500字"],
        ["trips[].subsidyDays", "List<SubsidyDaySaveDTO>", "N", "补助日历录入", "未传日期默认未选中"],
        ["allocations", "List<AllocationSaveDTO>", "N", "费用分摊", "提交时至少1条"],
    ]
    api_intro(doc, "A.10 创建草稿", "REIM_createDraft", "/api/reimbursements/drafts", "POST", "新增或首次自动保存报销单。")
    api_params(doc, save_payload_rows)
    api_output(doc, [["data", "ReimbursementDetailVO", "Y", "保存后的详情和生成单号", "状态草稿"]])
    api_intro(doc, "A.11 更新草稿", "REIM_saveDraft", "/api/reimbursements/{id}/draft", "PUT", "手工保存草稿或关闭页面自动保留。")
    api_params(doc, [["id", "Long", "Y", "报销单主键", "Path"]] + save_payload_rows)
    api_output(doc, [["data", "ReimbursementDetailVO", "Y", "更新后的整页详情", "按聚合快照返回"]])

    for title, name, path, scene, precondition in [
        ("A.12 提交报销单", "REIM_submit", "/api/reimbursements/{id}/submit", "提交完整报销单并变更状态。", "提交前必须通过全量校验。"),
        ("A.13 撤回报销单", "REIM_withdraw", "/api/reimbursements/{id}/withdraw", "将已完成单据撤回到草稿。", "仅已完成单据可撤回。"),
        ("A.14 作废报销单", "REIM_void", "/api/reimbursements/{id}/void", "本人作废未作废的报销单。", "仅本人可作废。"),
    ]:
        api_intro(doc, title, name, path, "POST", scene)
        api_params(doc, [["id", "Long", "Y", "报销单主键", "Path"], ["前置条件", "String", "Y", precondition, "-"]])
        api_output(doc, [
            ["id", "Long", "Y", "报销单主键", ""],
            ["reimNo", "String", "Y", "报销单号", ""],
            ["status", "Integer", "Y", "目标状态码", "0/1/2"],
            ["statusName", "String", "Y", "目标状态名", "草稿/已完成/已作废"],
        ])


def add_data_dictionary_appendix(doc: Document) -> None:
    heading(doc, "附录 B 数据字典与契约说明", 1)
    heading(doc, "B.1 补助规则", 2)
    add_table(doc, ["补助项", "标准", "规则"], [
        ["餐费补助", "一线100 / 二线80 / 三线50", "按到达城市 cityType 匹配。"],
        ["交通补助", "40 / 天", "所有城市固定标准。"],
        ["通讯补助", "40 / 天", "所有城市固定标准。"],
        ["实际金额", "0 到标准金额", "未勾选记0；勾选未传金额默认标准金额。"],
    ], [2200, 2700, 4460], font_size=8.4)
    heading(doc, "B.2 主要 DTO", 2)
    add_table(doc, ["DTO", "字段", "说明"], [
        ["LoginDTO", "username, password", "登录入参。"],
        ["ReimbursementDraftSaveDTO", "主单字段、trips、allocations", "草稿聚合保存入参。"],
        ["TripSaveDTO", "travelerId, departCityId, arriveCityId, departDate, arriveDate, tripDescription, subsidyDays", "行程及其补助明细。"],
        ["SubsidyDaySaveDTO", "subsidyDate 与餐补/交通/通讯选中状态及金额", "逐日补助输入。"],
        ["AllocationSaveDTO", "companyId, projectId, allocationRatio, allocationAmount", "费用分摊输入。"],
    ], [2600, 3600, 3160], font_size=8.1)
    heading(doc, "B.3 主要 VO", 2)
    add_table(doc, ["VO", "用途", "说明"], [
        ["LoginVO", "登录响应", "JWT 与有效期。"],
        ["ReimbursementListVO", "列表记录", "列表展示字段和状态。"],
        ["ReimbursementDetailVO", "详情响应", "主单、行程、补助、分摊整页数据。"],
        ["ReimbursementActionVO", "状态动作响应", "动作后单据 ID、单号、状态。"],
        ["BusinessTypeTreeVO", "业务类型树", "递归 children。"],
    ], [2600, 2600, 4160], font_size=8.3)
    heading(doc, "B.4 当前实现文件定位", 2)
    add_table(doc, ["主题", "文件"], [
        ["报销核心服务", r"src/main/java/com/kjd/travel/reimbursement/service/ReimbursementService.java"],
        ["报销接口", r"src/main/java/com/kjd/travel/reimbursement/controller/ReimbursementController.java"],
        ["基础数据接口", r"src/main/java/com/kjd/travel/masterdata/controller/MasterDataController.java"],
        ["登录接口", r"src/main/java/com/kjd/travel/auth/controller/AuthController.java"],
        ["建表 SQL", r"src/main/resources/db/schema.sql"],
        ["测试数据", r"src/main/resources/db/data.sql"],
        ["核心测试", r"src/test/java/com/kjd/travel/reimbursement/ReimbursementServiceTests.java"],
    ], [2500, 6860], font_size=8.0)


def build_document() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_scope_and_terms(doc)
    add_requirements_and_functions(doc)
    add_function_design(doc)
    add_technical_design(doc)
    add_key_technical_points(doc)
    add_database_design(doc)
    add_wbs_and_verification(doc)
    add_api_appendix(doc)
    add_data_dictionary_appendix(doc)
    doc.save(OUT_FILE)
    return OUT_FILE


if __name__ == "__main__":
    print(build_document())
