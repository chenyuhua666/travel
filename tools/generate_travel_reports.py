from __future__ import annotations

import csv
import json
import os
import subprocess
from collections import defaultdict
from dataclasses import dataclass
from io import StringIO
from pathlib import Path
from typing import Iterable, Sequence

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


WORKSPACE = Path(r"E:\codelocation\IDEA\travel")
FRONTEND = Path(r"E:\codelocation\vue\travel\travel")
TEMPLATE = Path(r"E:\qq下载路径\差旅报销系统_详细设计WBS.docx")
OUT_DIR = Path(r"E:\qq下载路径")
ASSET_DIR = WORKSPACE / "outputs" / "report-assets"
WBS_TASKS_FILE = WORKSPACE / "tools" / "wbs_tasks.json"

DETAIL_DOCX = OUT_DIR / "差旅报销系统_详细设计报告_更新版.docx"
WBS_DOCX = OUT_DIR / "差旅报销系统_WBS报告_更新版.docx"

CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(30, 39, 52)
MUTED = RGBColor(95, 105, 118)
LIGHT_FILL = "F2F4F7"
BLUE_FILL = "E8EEF5"
CALLOUT_FILL = "F4F6F9"
WARN_FILL = "FFF7E6"

FONT_CN = "Microsoft YaHei"
FONT_WEST = "Calibri"


@dataclass
class DbInfo:
    tables: list[str]
    row_counts: dict[str, int]
    columns: dict[str, list[dict[str, str]]]
    indexes: dict[str, list[dict[str, str]]]
    fks: dict[str, list[dict[str, str]]]
    status_counts: list[dict[str, str]]


TABLE_DESCRIPTIONS = {
    "sys_company": "公司基础表",
    "sys_department": "部门基础表",
    "sys_employee": "员工基础表",
    "sys_user": "登录用户表",
    "biz_business_type": "业务类型树",
    "biz_city": "城市及补助等级",
    "biz_project": "费用归集项目",
    "fk_reim_main": "报销单主表",
    "fk_reim_trip": "补录行程表",
    "fk_reim_subsidy": "补助汇总表",
    "fk_reim_subsidy_day": "补助日历明细表",
    "fk_reim_allocation": "费用归属分摊表",
}

FIELD_DESCRIPTIONS = {
    "id": "主键ID",
    "created_at": "创建时间",
    "creation_time": "创建时间",
    "update_time": "更新时间",
    "company_no": "公司编号",
    "company_name": "公司名称",
    "department_no": "部门编号",
    "department_name": "部门名称",
    "employee_no": "员工工号",
    "employee_name": "员工姓名",
    "username": "登录账号",
    "password_hash": "PBKDF2密码哈希",
    "roles": "角色集合，逗号分隔",
    "enabled": "账号启用标识",
    "business_type_no": "业务类型编号",
    "business_type_name": "业务类型名称",
    "parent_id": "父业务类型ID",
    "leaf_flag": "是否叶子节点",
    "city_no": "城市编号",
    "city_name": "城市名称",
    "city_type": "城市等级，1一线/2二线/3三线",
    "project_no": "项目编号",
    "project_name": "项目名称",
    "reim_no": "报销单号",
    "reimbursement_title": "报销标题",
    "reimburser_id": "报销人ID",
    "reimburser_no": "报销人工号快照",
    "reimburser_name": "报销人姓名快照",
    "reim_department_id": "报销部门ID",
    "reim_department_no": "报销部门编号快照",
    "reim_department_name": "报销部门名称快照",
    "reim_company_id": "费用归属公司ID",
    "reim_company_no": "费用归属公司编号快照",
    "reim_company_name": "费用归属公司名称快照",
    "business_type_id": "业务类型ID",
    "business_trip_reason": "出差事由",
    "subsidy_total": "补助总金额",
    "meal_allowance": "餐费补助合计",
    "transportation_allowance": "交通补助合计",
    "phone_allowance": "通讯补助合计",
    "remarks": "备注",
    "status": "单据状态",
    "owner_user_id": "单据所有人用户ID",
    "main_id": "报销主表ID",
    "traveler_id": "出行人员ID",
    "traveler_no": "出行人员工号快照",
    "traveler_name": "出行人员姓名快照",
    "depart_city_id": "出发城市ID",
    "depart_city_no": "出发城市编号快照",
    "depart_city_name": "出发城市名称快照",
    "arrive_city_id": "到达城市ID",
    "arrive_city_no": "到达城市编号快照",
    "arrive_city_name": "到达城市名称快照",
    "depart_date": "出发日期",
    "arrive_date": "到达日期",
    "trip_description": "行程说明",
    "trip_id": "行程ID",
    "subsidy_days": "补助天数",
    "apply_amount": "申请标准金额",
    "subsidy_amount": "实际补助金额",
    "meal_amount": "餐费补助金额",
    "transportation_amount": "交通补助金额",
    "phone_amount": "通讯补助金额",
    "subsidy_id": "补助汇总ID",
    "subsidy_date": "补助日期",
    "weekday_name": "星期名称",
    "meal_standard_amount": "餐费标准金额",
    "transportation_standard_amount": "交通标准金额",
    "phone_standard_amount": "通讯标准金额",
    "meal_selected": "餐费是否勾选",
    "transportation_selected": "交通是否勾选",
    "phone_selected": "通讯是否勾选",
    "allocation_ratio": "分摊比例，0到1",
    "allocation_amount": "分摊金额",
    "row_order": "分摊行顺序",
}


def mysql_query(sql: str) -> list[dict[str, str]]:
    password = os.environ.get("TRAVEL_DB_PASSWORD", "")
    username = os.environ.get("TRAVEL_DB_USERNAME", "root")
    database = os.environ.get("TRAVEL_DB_NAME", "travel")
    mysql_exe = os.environ.get("MYSQL_EXE", "mysql")
    args = [
        mysql_exe,
        "--batch",
        "--raw",
        "--default-character-set=utf8mb4",
        f"-u{username}",
    ]
    if password:
        args.append(f"-p{password}")
    args.extend([database, "-e", sql])
    completed = subprocess.run(args, check=True, capture_output=True, text=True, encoding="utf-8")
    reader = csv.DictReader(StringIO(completed.stdout), delimiter="\t")
    return [dict(row) for row in reader]


def load_db_info() -> DbInfo:
    tables = [
        row["TABLE_NAME"]
        for row in mysql_query(
            "SELECT TABLE_NAME FROM information_schema.TABLES "
            "WHERE TABLE_SCHEMA=DATABASE() ORDER BY TABLE_NAME"
        )
    ]
    row_counts: dict[str, int] = {}
    for table in tables:
        count_rows = mysql_query(f"SELECT COUNT(*) AS row_count FROM `{table}`")
        row_counts[table] = int(count_rows[0]["row_count"]) if count_rows else 0

    columns: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in mysql_query(
        "SELECT TABLE_NAME,COLUMN_NAME,COLUMN_TYPE,IS_NULLABLE,COLUMN_DEFAULT,COLUMN_KEY,EXTRA "
        "FROM information_schema.COLUMNS WHERE TABLE_SCHEMA=DATABASE() "
        "ORDER BY TABLE_NAME,ORDINAL_POSITION"
    ):
        columns[row["TABLE_NAME"]].append(row)

    indexes: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in mysql_query(
        "SELECT TABLE_NAME,INDEX_NAME,NON_UNIQUE,"
        "GROUP_CONCAT(COLUMN_NAME ORDER BY SEQ_IN_INDEX) AS columns "
        "FROM information_schema.STATISTICS WHERE TABLE_SCHEMA=DATABASE() "
        "GROUP BY TABLE_NAME,INDEX_NAME,NON_UNIQUE ORDER BY TABLE_NAME,INDEX_NAME"
    ):
        indexes[row["TABLE_NAME"]].append(row)

    fks: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in mysql_query(
        "SELECT TABLE_NAME,COLUMN_NAME,CONSTRAINT_NAME,REFERENCED_TABLE_NAME,REFERENCED_COLUMN_NAME "
        "FROM information_schema.KEY_COLUMN_USAGE "
        "WHERE TABLE_SCHEMA=DATABASE() AND REFERENCED_TABLE_NAME IS NOT NULL "
        "ORDER BY TABLE_NAME,COLUMN_NAME"
    ):
        fks[row["TABLE_NAME"]].append(row)

    status_counts = mysql_query(
        "SELECT status, COUNT(*) AS row_count FROM fk_reim_main GROUP BY status ORDER BY status"
    )
    return DbInfo(tables=tables, row_counts=row_counts, columns=columns, indexes=indexes, fks=fks, status_counts=status_counts)


def clear_body(doc: Document) -> None:
    body = doc._body._element
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_run(run, size: float | None = None, bold: bool | None = None, color: RGBColor | None = None) -> None:
    run.font.name = FONT_WEST
    if run._element.rPr is None:
        run._element.get_or_add_rPr()
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT_WEST)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_WEST)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    run._r.append(fld_begin)

    run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    run._r.append(instr)

    run = paragraph.add_run()
    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")
    run._r.append(fld_separate)

    if placeholder:
        set_run(paragraph.add_run(placeholder), 8.5, False, MUTED)

    run = paragraph.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_end)


def enable_update_fields(doc: Document) -> None:
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def configure_document(doc: Document, running_title: str) -> None:
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_WEST
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in (
        ("Title", 24, INK, 0, 8),
        ("Subtitle", 13, DARK_BLUE, 0, 14),
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        if style_name in styles:
            style = styles[style_name]
            style.font.name = FONT_WEST
            style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
            style.font.size = Pt(size)
            style.font.bold = style_name.startswith("Heading") or style_name == "Title"
            style.font.color.rgb = color
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)
            style.paragraph_format.line_spacing = 1.15

    for style_name in ("List Bullet", "List Number"):
        if style_name in styles:
            style = styles[style_name]
            style.font.name = FONT_WEST
            style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
            style.font.size = Pt(10.5)
            style.paragraph_format.space_after = Pt(4)
            style.paragraph_format.line_spacing = 1.167

    header = section.header
    hp = header.paragraphs[0]
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(hp.add_run(running_title), 9, False, MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(fp.add_run("内部设计文档 | 数据库只读核对 | 2026-05-25 | 第 "), 8.5, False, MUTED)
    add_field(fp, "PAGE", "1")
    set_run(fp.add_run(" 页 / 共 "), 8.5, False, MUTED)
    add_field(fp, "NUMPAGES", "1")
    set_run(fp.add_run(" 页"), 8.5, False, MUTED)
    enable_update_fields(doc)


def paragraph(doc: Document, text: str, style: str | None = None, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run(r1, 11, True, INK)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run(r2, 11, False, INK)
    else:
        set_run(p.add_run(text), 11, False, INK)


def bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.08)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    set_run(p.add_run(text), 10.5, False, INK)


def number_item(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.08)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    set_run(p.add_run(text), 10.5, False, INK)


def heading(doc: Document, text: str, level: int = 1) -> None:
    try:
        p = doc.add_paragraph(style=f"Heading {level}")
    except KeyError:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8 if level >= 3 else 12)
        p.paragraph_format.space_after = Pt(4 if level >= 3 else 6)
        p.paragraph_format.line_spacing = 1.15
    set_run(p.add_run(text), 12 if level >= 3 else None, True, BLUE if level < 3 else DARK_BLUE)


def page_break(doc: Document) -> None:
    doc.add_page_break()


def set_cell_margins(cell, top: int = 80, bottom: int = 80, start: int = 120, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: object, bold: bool = False, size: float = 8.5, color: RGBColor = INK) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.06
    set_run(p.add_run("" if text is None else str(text)), size, bold, color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths: Sequence[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    grid = tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[min(index, len(widths) - 1)]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            set_cell_margins(cell)


def add_table(
    doc: Document,
    headers: Sequence[str],
    rows: Iterable[Sequence[object]],
    widths: Sequence[int],
    font_size: float = 8.3,
    header_fill: str = LIGHT_FILL,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_text(cell, header, True, font_size, DARK_BLUE)
        set_cell_shading(cell, header_fill)
    set_repeat_header(table.rows[0])
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            set_cell_text(cell, value, False, font_size, INK)
    set_table_geometry(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)


def add_kv_table(doc: Document, rows: Sequence[tuple[str, str]]) -> None:
    add_table(doc, ["项目", "内容"], rows, [2200, 7160], font_size=8.8)


def callout(doc: Document, title: str, body: str, fill: str = CALLOUT_FILL) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_repeat_header(table.rows[0])
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, fill)
    set_cell_margins(cell, 120, 120, 160, 160)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    set_run(p.add_run(title), 10.2, True, DARK_BLUE)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.12
    set_run(p2.add_run(body), 9.4, False, INK)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def cover(doc: Document, title: str, subtitle: str, meta: Sequence[tuple[str, str]]) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(38)
    p.paragraph_format.space_after = Pt(4)
    set_run(p.add_run("差旅报销系统"), 15, False, DARK_BLUE)
    p = doc.add_paragraph(style="Title")
    set_run(p.add_run(title), 26, True, INK)
    p = doc.add_paragraph()
    set_run(p.add_run(subtitle), 13, False, DARK_BLUE)
    add_kv_table(doc, meta)
    callout(
        doc,
        "文档依据",
        "本文档以《差旅报销系统_详细设计WBS.docx》为模板，结合《差旅报销单概要设计》、前后端源码以及 3306 travel 数据库只读元数据核对结果编制。",
    )


def toc(doc: Document, items: Sequence[str]) -> None:
    heading(doc, "目录", 1)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    add_field(p, r'TOC \o "1-3" \h \z \u', "目录将在打开文档时自动更新页码")


def font(size: int, bold: bool = False):
    path = Path(r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc")
    return ImageFont.truetype(str(path), size=size)


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font_obj: ImageFont.ImageFont, max_width: int) -> list[str]:
    lines: list[str] = []
    current = ""
    for char in text:
        test = current + char
        if draw.textbbox((0, 0), test, font=font_obj)[2] <= max_width:
            current = test
        else:
            if current:
                lines.append(current)
            current = char
    if current:
        lines.append(current)
    return lines


def draw_centered_box(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    fill: str,
    outline: str = "#6B7C93",
    text_fill: str = "#1E2734",
    bold: bool = False,
) -> None:
    draw.rounded_rectangle(box, radius=18, fill=fill, outline=outline, width=2)
    f = font(30, bold)
    lines = wrap_text(draw, text, f, box[2] - box[0] - 34)
    line_h = 40
    y = box[1] + ((box[3] - box[1]) - line_h * len(lines)) // 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=f)
        x = box[0] + ((box[2] - box[0]) - (bbox[2] - bbox[0])) // 2
        draw.text((x, y), line, font=f, fill=text_fill)
        y += line_h


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], fill: str = "#4A617A") -> None:
    draw.line((start, end), fill=fill, width=4)
    ex, ey = end
    sx, sy = start
    if abs(ex - sx) >= abs(ey - sy):
        direction = 1 if ex > sx else -1
        pts = [(ex, ey), (ex - 16 * direction, ey - 10), (ex - 16 * direction, ey + 10)]
    else:
        direction = 1 if ey > sy else -1
        pts = [(ex, ey), (ex - 10, ey - 16 * direction), (ex + 10, ey - 16 * direction)]
    draw.polygon(pts, fill=fill)


def diagram_base(title: str, size: tuple[int, int] = (1600, 900)) -> tuple[Image.Image, ImageDraw.ImageDraw]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, size[0], 88), fill="#E8EEF5")
    draw.text((44, 22), title, font=font(34, True), fill="#1F4D78")
    return image, draw


def save_architecture_diagram() -> Path:
    path = ASSET_DIR / "architecture.png"
    image, draw = diagram_base("系统框架图")
    boxes = [
        ((70, 170, 350, 310), "浏览器 / 用户"),
        ((450, 150, 790, 330), "Vue 3 前端\nRouter / Pinia / Axios\nElement Plus"),
        ((910, 130, 1245, 275), "Spring Boot 4\nController 层"),
        ((910, 340, 1245, 500), "Service 层\n事务 / 校验 / 计算"),
        ((910, 570, 1245, 720), "MyBatis-Plus\nMapper / Entity"),
        ((1340, 430, 1540, 630), "MySQL 8\ntravel库"),
        ((450, 440, 790, 610), "JWT 安全链\n登录放行 / 其余鉴权"),
    ]
    for box, text in boxes:
        draw_centered_box(draw, box, text, "#F7FAFD", bold=True)
    arrow(draw, (350, 240), (450, 240))
    arrow(draw, (790, 240), (910, 205))
    arrow(draw, (1075, 275), (1075, 340))
    arrow(draw, (1075, 500), (1075, 570))
    arrow(draw, (1245, 645), (1340, 530))
    arrow(draw, (620, 440), (620, 330))
    arrow(draw, (790, 520), (910, 205))
    draw.text((80, 790), "核心调用链：页面操作 -> Axios请求 -> JWT过滤器 -> Controller -> Service -> Mapper -> MySQL", font=font(24), fill="#5F6976")
    image.save(path)
    return path


def save_business_flow_diagram() -> Path:
    path = ASSET_DIR / "business_flow.png"
    image, draw = diagram_base("业务流程图")
    boxes = [
        ((80, 170, 320, 270), "登录获取Token"),
        ((400, 170, 640, 270), "查询报销列表"),
        ((720, 170, 960, 270), "新增/编辑草稿"),
        ((1040, 170, 1280, 270), "填写基础信息"),
        ((80, 390, 320, 490), "补录行程"),
        ((400, 390, 640, 490), "生成补助日历"),
        ((720, 390, 960, 490), "费用归属分摊"),
        ((1040, 390, 1280, 490), "保存草稿"),
        ((400, 620, 640, 720), "提交校验"),
        ((720, 620, 960, 720), "审批中"),
        ((1040, 590, 1280, 690), "审批通过"),
        ((1040, 730, 1280, 830), "撤回/作废"),
    ]
    for box, text in boxes:
        draw_centered_box(draw, box, text, "#FFFFFF", "#6B7C93", bold=True)
    for a, b in [
        ((320, 220), (400, 220)),
        ((640, 220), (720, 220)),
        ((960, 220), (1040, 220)),
        ((1160, 270), (200, 390)),
        ((320, 440), (400, 440)),
        ((640, 440), (720, 440)),
        ((960, 440), (1040, 440)),
        ((1160, 490), (520, 620)),
        ((640, 670), (720, 670)),
        ((960, 670), (1040, 640)),
        ((840, 720), (1040, 780)),
    ]:
        arrow(draw, a, b)
    image.save(path)
    return path


def save_sequence_diagram() -> Path:
    path = ASSET_DIR / "sequence_save_submit.png"
    image, draw = diagram_base("程序时序图：保存草稿与提交")
    participants = ["用户", "Vue页面", "Axios API", "JWT过滤器", "Controller", "Service", "Mapper", "MySQL"]
    xs = [90, 285, 480, 690, 890, 1080, 1265, 1460]
    top = 150
    bottom = 790
    for x, p in zip(xs, participants):
        draw_centered_box(draw, (x - 72, top, x + 72, top + 70), p, "#F7FAFD", bold=True)
        draw.line((x, top + 70, x, bottom), fill="#B8C2CC", width=2)
    steps = [
        (0, 1, "点击保存/提交", 260),
        (1, 2, "组装DraftPayload", 310),
        (2, 3, "携带Bearer Token", 360),
        (3, 4, "认证通过注入CurrentUser", 410),
        (4, 5, "createDraft/saveDraft/submit", 460),
        (5, 6, "事务内校验、计算、替换子表", 510),
        (6, 7, "insert/update/select", 560),
        (7, 6, "返回持久化结果", 610),
        (6, 5, "实体转VO", 660),
        (5, 4, "Result<T>", 700),
        (4, 2, "统一响应", 740),
        (2, 1, "刷新页面状态", 780),
    ]
    for src, dst, label, y in steps:
        arrow(draw, (xs[src] + 12, y), (xs[dst] - 12 if dst > src else xs[dst] + 12, y), "#4A617A")
        draw.text((min(xs[src], xs[dst]) + 22, y - 28), label, font=font(20), fill="#1E2734")
    image.save(path)
    return path


def save_wbs_diagram() -> Path:
    path = ASSET_DIR / "wbs_tree.png"
    image, draw = diagram_base("WBS框架图")
    root = (590, 130, 1010, 220)
    draw_centered_box(draw, root, "差旅报销系统交付", "#E8EEF5", bold=True)
    level1 = [
        ((60, 340, 300, 430), "1 需求与范围"),
        ((350, 340, 590, 430), "2 详细设计"),
        ((640, 340, 880, 430), "3 后端实现"),
        ((930, 340, 1170, 430), "4 前端实现"),
        ((1220, 340, 1460, 430), "5 测试交付"),
    ]
    for box, text in level1:
        draw_centered_box(draw, box, text, "#FFFFFF", bold=True)
        arrow(draw, ((root[0] + root[2]) // 2, root[3]), ((box[0] + box[2]) // 2, box[1]))
    details = [
        ((60, 560, 300, 690), "流程梳理\n状态定义\n原型映射"),
        ((350, 560, 590, 690), "接口文档\n数据库表\n时序与框架"),
        ((640, 560, 880, 690), "鉴权\n业务服务\nMySQL持久化"),
        ((930, 560, 1170, 690), "列表\n表单\n补助日历"),
        ((1220, 560, 1460, 690), "自测\n联调\n文档归档"),
    ]
    for box, text in details:
        draw_centered_box(draw, box, text, "#F7FAFD")
        cx = (box[0] + box[2]) // 2
        arrow(draw, (cx, 430), (cx, 560), "#7A8FA6")
    image.save(path)
    return path


def add_image(doc: Document, path: Path, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    shape = run.add_picture(str(path), width=Inches(6.25))
    shape._inline.docPr.set("descr", caption)
    shape._inline.docPr.set("title", caption)
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(c.add_run(caption), 9, False, MUTED)


def revision_table(doc: Document) -> None:
    heading(doc, "修订记录", 1)
    add_table(
        doc,
        ["日期", "版本", "修改章节", "修改说明", "作者"],
        [["2026-05-25", "V1.0", "全量", "依据模板、代码、数据库和项目资料生成正式报告", "Codex"]],
        [1500, 900, 1600, 4360, 1000],
        8.8,
    )


def add_nonfunctional(doc: Document) -> None:
    heading(doc, "1 非功能性要求", 1)
    add_table(
        doc,
        ["维度", "设计要求", "当前落点"],
        [
            ["可靠性", "草稿保存、提交、撤回、审批、作废等关键动作应保持主子表一致。", "ReimbursementService 关键写操作使用 @Transactional；保存时先校验再替换行程、补助、分摊子表。"],
            ["高性能", "列表查询支持分页、条件过滤和常用索引，基础数据轻量返回。", "MyBatis-Plus 分页最大100条；主表状态、创建时间、报销人、公司、所有人等字段建索引。"],
            ["可维护性", "分层清晰，DTO/VO/Entity 隔离，接口与页面契约稳定。", "controller/service/mapper/entity/dto/vo/config 分层；前端 api/types/stores/views 分层。"],
            ["安全性", "除登录外均需认证，用户只能访问或操作本人单据。", "Spring Security 无状态 JWT；owner_user_id 做数据权限；密码采用 PBKDF2 校验。"],
            ["可扩展性", "审批流、票据识别、发票查验可后续接入。", "当前通过 APPROVING/APPROVED 状态预留审批节点，第三方接口尚未接入。"],
        ],
        [1500, 3920, 3940],
        8.3,
    )


def add_terms(doc: Document) -> None:
    heading(doc, "2 术语定义", 1)
    add_table(
        doc,
        ["术语", "定义"],
        [
            ["报销单主单", "承载报销标题、报销人、部门、费用归属公司、业务类型、事由、补助合计、备注和状态的聚合根。"],
            ["补录行程", "用户手工录入的出行人、出发/到达城市、出发/到达日期和行程说明。"],
            ["补助日历", "按行程日期逐日生成的餐费、交通、通讯补助勾选与金额明细。"],
            ["费用分摊", "将补助总金额按费用归属公司、项目、分摊比例和分摊金额拆分。"],
            ["草稿", "状态码0，允许编辑、删除、提交。"],
            ["审批中", "状态码3，提交后的待审批状态，允许撤回、审批通过或作废。"],
            ["审批通过", "状态码1，审批完成后的状态。"],
            ["已作废", "状态码2，终止状态。"],
            ["基础数据", "公司、部门、员工、业务类型、城市、项目等页面控件数据。"],
            ["Result<T>", "后端统一响应结构，成功 code=0、message=success、data 为业务数据。"],
        ],
        [1900, 7460],
        8.7,
    )


def add_requirements(doc: Document, diagrams: dict[str, Path], db: DbInfo) -> None:
    heading(doc, "3 功能性需求描述", 1)
    heading(doc, "3.1 需求概述", 2)
    paragraph(
        doc,
        "本系统实现费控云差旅报销模块，覆盖登录、基础数据加载、报销单列表查询、新增、编辑、保存草稿、提交、撤回、审批通过、作废、复制、删除草稿，以及补录行程、补助日历、费用合计和费用分摊。前端提供可操作页面，后端提供 REST API，数据库使用 MySQL travel 库持久化。",
    )
    heading(doc, "3.2 用户角色", 2)
    add_table(
        doc,
        ["角色", "本期权限", "说明"],
        [
            ["员工/提单人", "登录、查询本人单据、新增、编辑草稿、保存草稿、提交、撤回、复制、删除草稿、作废本人单据。", "当前代码按 owner_user_id 控制本人数据。"],
            ["管理员", "可扩展为全量查询和基础数据维护。", "CurrentUser.hasRole(\"ADMIN\") 已用于列表可见范围判断。"],
            ["审批人", "审批中单据可通过或作废。", "当前 approve 接口存在，未单独建审批人表或流程引擎。"],
        ],
        [1900, 4900, 2560],
        8.4,
    )
    heading(doc, "3.3 单据状态", 2)
    status_map = {"0": "草稿", "1": "审批通过", "2": "已作废", "3": "审批中"}
    status_rows = [[code, name, next((r["row_count"] for r in db.status_counts if r["status"] == code), "0")] for code, name in status_map.items()]
    add_table(doc, ["状态码", "状态名称", "当前 travel 库样例数量"], status_rows, [1800, 3200, 4360], 8.6)
    heading(doc, "3.4 业务流程", 2)
    add_image(doc, diagrams["business"], "图3-1 差旅报销业务流程")
    heading(doc, "3.5 功能范围", 2)
    add_table(
        doc,
        ["模块", "功能点", "对应前后端实现"],
        [
            ["认证", "登录、Token保存、请求鉴权、401自动回登录页", "AuthController/AuthService/JwtAuthenticationFilter；src/stores/auth.ts、src/api/http.ts"],
            ["基础数据", "公司、部门、员工、业务类型树、城市、项目下拉", "MasterDataController/MasterDataService；src/stores/master.ts"],
            ["报销列表", "多条件查询、分页、新增入口、编辑、审批、复制、删除", "ReimbursementController.page；ReimbursementListView.vue"],
            ["报销表单", "基础信息、补录行程、补助信息、费用合计、分摊、备注", "ReimbursementFormView.vue；ReimbursementService.saveDraft"],
            ["状态流转", "提交、撤回、审批通过、作废", "submit/withdraw/approve/void 接口"],
            ["数据持久化", "主表、行程、补助汇总、补助日历、费用分摊", "fk_reim_main、fk_reim_trip、fk_reim_subsidy、fk_reim_subsidy_day、fk_reim_allocation"],
        ],
        [1600, 3300, 4460],
        8.0,
    )


def add_function_design(doc: Document, diagrams: dict[str, Path]) -> None:
    heading(doc, "4 功能详细设计", 1)
    heading(doc, "4.1 登录与鉴权", 2)
    paragraph(doc, "登录接口接收用户名和密码，AuthService 按 username 查询 sys_user，校验 enabled 与 PBKDF2 密码哈希，通过后生成 HS256 JWT。前端将 token 写入 localStorage，Axios 拦截器自动添加 Authorization: Bearer {token}。")
    add_table(
        doc,
        ["处理环节", "逻辑", "异常处置"],
        [
            ["登录", "校验用户名、密码、账号启用状态，生成 tokenType=Bearer 的 JWT。", "用户名或密码错误返回 code=401。"],
            ["请求鉴权", "JwtAuthenticationFilter 解析 token，生成 CurrentUser 注入 Spring Security 上下文。", "令牌格式错误、签名无效、过期或解析失败返回身份认证失败。"],
            ["数据权限", "列表非管理员按 owner_user_id 过滤；详情/编辑/状态操作需本人单据。", "非本人访问返回 code=403。"],
        ],
        [1800, 5200, 2360],
        8.1,
    )

    heading(doc, "4.2 基础数据", 2)
    paragraph(doc, "基础数据用于页面下拉控件和业务类型树。前端 master store 首次进入页面时并发加载公司、部门、员工、业务类型、城市、项目，后续复用缓存。业务类型按 parent_id 递归组装树形结构，leaf_flag 用于区分可选叶子。")
    add_table(
        doc,
        ["数据项", "接口", "页面用途"],
        [
            ["公司", "GET /api/master/companies", "费用归属公司、分摊费用归属。"],
            ["部门", "GET /api/master/departments", "报销部门查询与表单选择。"],
            ["员工", "GET /api/master/employees", "报销人、出行人选择。"],
            ["业务类型", "GET /api/master/business-types/tree", "列表查询与报销单业务类型树形选择。"],
            ["城市", "GET /api/master/cities", "行程出发/到达城市与补助标准计算。"],
            ["项目", "GET /api/master/projects", "费用分摊项目选择。"],
        ],
        [1600, 3100, 4660],
        8.3,
    )

    heading(doc, "4.3 报销单列表", 2)
    for text in [
        "查询条件包括报销单号、标题、事由、费用归属公司、报销部门、报销人、业务类型；分页参数 current、size 默认 1/10，后端限制 size 最大 100。",
        "列表展示操作、报销单号、单据状态、单据类型、报销人、报销部门、费用归属公司、业务类型、报销标题、报销事由、补助金额、创建时间。",
        "草稿可编辑和删除；审批中可审批通过或作废；任意可见单据可复制为新草稿；点击单号进入表单页。",
    ]:
        bullet(doc, text)

    heading(doc, "4.4 报销单表单与基础信息", 2)
    paragraph(doc, "表单页面分为基础信息、补录行程、补助信息、费用合计、费用归属及分摊、备注信息。单据头部固定展示单号、状态、标题和提单日期；状态不是草稿时前端进入只读模式。")
    add_table(
        doc,
        ["字段", "校验", "存储/展示"],
        [
            ["报销标题", "必填，最长500字", "fk_reim_main.reimbursement_title"],
            ["报销人", "必选", "保存员工ID、工号、姓名快照"],
            ["报销部门", "必选", "保存部门ID、编号、名称快照"],
            ["费用归属公司", "必选", "保存公司ID、编号、名称快照，并可带入分摊首行"],
            ["业务类型", "必选", "保存业务类型ID、编号、名称快照"],
            ["出差事由", "必填，最长500字", "fk_reim_main.business_trip_reason"],
            ["备注", "最长1000字", "fk_reim_main.remarks"],
        ],
        [1800, 2600, 4960],
        8.2,
    )

    heading(doc, "4.5 补录行程", 2)
    paragraph(doc, "补录行程弹窗包含出行人、出发城市、到达城市、出发/到达日期、行程说明。前端先做必填、日期范围和单内同人日期重叠校验；后端在保存时再次校验，到达日期不可早于出发日期且不可晚于当前日期。")
    add_table(
        doc,
        ["动作", "前端逻辑", "后端逻辑"],
        [
            ["新增", "填写弹窗后生成 TripRow，并按到达城市生成默认补助日历。", "保存时写入 fk_reim_trip，并生成 fk_reim_subsidy/fk_reim_subsidy_day。"],
            ["编辑", "回显当前行，若日期或到达城市未变则保留补助选择。", "聚合保存时先删除旧子表，再重建最新快照。"],
            ["复制", "复制当前行进入弹窗，保存后生成新行。", "copy 接口可复制整张单据为新草稿。"],
            ["删除", "二次确认后移除行程并重算分摊。", "保存后子表按最新请求替换。"],
        ],
        [1500, 4000, 3860],
        8.1,
    )

    heading(doc, "4.6 补助信息与补助日历", 2)
    paragraph(doc, "补助标准按到达城市 city_type 计算：一线城市餐费100元/天，二线80元/天，三线50元/天；交通补助和通讯补助统一40元/天。用户可按行、按列或全选勾选补助项，未勾选金额记0，勾选金额必须在0到标准金额之间。")
    add_table(
        doc,
        ["补助项", "标准", "计算逻辑"],
        [
            ["餐费补助", "城市等级：100/80/50 元/天", "选中时默认标准金额，可人工下调，不能超过标准。"],
            ["交通补助", "40 元/天", "所有城市固定标准。"],
            ["通讯补助", "40 元/天", "所有城市固定标准。"],
            ["申请金额", "标准金额合计", "按已勾选补助项的标准金额求和。"],
            ["补助金额", "实际金额合计", "按用户实际录入金额求和，回写主表合计。"],
        ],
        [1800, 2500, 5060],
        8.3,
    )

    heading(doc, "4.7 费用合计与费用分摊", 2)
    paragraph(doc, "费用合计展示补助总金额、餐费、交通、通讯分项合计。分摊明细包含费用归属公司、项目、分摊比例、分摊金额和顺序号。前端第二行起可编辑比例，首行自动承接剩余比例；均摊按钮按行数平分比例并重算金额。")
    add_table(
        doc,
        ["校验点", "规则"],
        [
            ["分摊行", "提交时至少一条，且每条有效分摊必须有费用归属公司。"],
            ["比例范围", "allocationRatio 需在0到1之间，提交时合计必须等于1。"],
            ["金额范围", "金额不可为负，保留2位小数。"],
            ["金额合计", "提交时分摊金额合计必须等于主表 subsidy_total。"],
        ],
        [2200, 7160],
        8.5,
    )

    heading(doc, "4.8 保存草稿与提交", 2)
    add_image(doc, diagrams["sequence"], "图4-1 保存草稿与提交程序时序图")
    for text in [
        "createDraft/saveDraft 使用同一聚合保存逻辑：填充主表快照字段，插入或更新主表，删除旧子表，按请求重建行程、补助、补助日历和分摊。",
        "新增草稿插入后生成单号，当前规则为 SY + 10位主键序号。",
        "submit 仅允许草稿状态，提交前执行主表必填、行程非空、行程重叠、分摊非空、比例合计100%、金额合计等于补助总额校验；成功后状态变为审批中。",
    ]:
        bullet(doc, text)

    heading(doc, "4.9 状态操作", 2)
    add_table(
        doc,
        ["接口", "前置条件", "结果"],
        [
            ["POST /api/reimbursements/{id}/withdraw", "本人单据，状态为审批中", "状态回到草稿。"],
            ["POST /api/reimbursements/{id}/approve", "本人单据，状态为审批中", "状态变为审批通过。"],
            ["POST /api/reimbursements/{id}/void", "本人单据，且未作废", "状态变为已作废。"],
            ["POST /api/reimbursements/{id}/copy", "可见单据", "复制主单、行程、补助日历、分摊为新草稿。"],
            ["DELETE /api/reimbursements/{id}", "本人草稿单据", "删除主表，子表由外键级联删除。"],
        ],
        [3300, 3000, 3060],
        8.0,
    )

    heading(doc, "4.10 异常处置", 2)
    paragraph(doc, "业务异常由 BusinessException 携带 code 与 message，统一返回 Result.failure；参数校验、绑定错误、认证失败、权限不足和系统异常分别映射为 400、401、403、500。前端 Axios 拦截器展示错误消息，401 时清除 token 并跳转登录页。")


def add_technical_design(doc: Document, diagrams: dict[str, Path]) -> None:
    heading(doc, "5 技术实现设计", 1)
    heading(doc, "5.1 系统结构设计", 2)
    add_image(doc, diagrams["architecture"], "图5-1 系统框架图")
    add_table(
        doc,
        ["层次", "主要组件", "职责"],
        [
            ["前端表现层", "Vue 3、Element Plus、Router、Pinia", "页面渲染、路由守卫、基础数据缓存、表单交互和前端校验。"],
            ["前端接口层", "src/api/http.ts、auth.ts、master.ts、reimbursements.ts", "Axios实例、Token注入、统一响应剥离和错误提示。"],
            ["后端控制层", "AuthController、MasterDataController、ReimbursementController", "REST入口、参数接收、统一Result响应。"],
            ["后端业务层", "AuthService、MasterDataService、ReimbursementService", "认证、树构建、报销聚合保存、补助计算、状态流转。"],
            ["持久化层", "MyBatis-Plus Mapper、Entity", "表映射、分页查询、主子表持久化。"],
            ["数据库层", "MySQL travel", "基础数据、用户、报销主表、行程、补助、补助日历、分摊。"],
        ],
        [1600, 3300, 4460],
        8.0,
    )
    heading(doc, "5.2 接口及核心类设计", 2)
    add_table(
        doc,
        ["核心类/文件", "职责"],
        [
            ["SecurityConfig", "禁用CSRF、启用无状态会话、登录接口放行，其余接口认证。"],
            ["JwtTokenService", "创建和解析HS256 JWT，校验签名与过期时间。"],
            ["PasswordHasher", "PBKDF2WithHmacSHA256 密码校验。"],
            ["ReimbursementService", "报销核心聚合服务，处理保存、提交、审批、撤回、作废、复制、删除。"],
            ["ReimbursementDraftSaveDTO", "草稿聚合保存入参，包含主表字段、行程、补助日历、分摊。"],
            ["ReimbursementDetailVO", "详情出参，包含主单、行程、补助汇总/明细、分摊。"],
            ["ReimbursementStatus", "状态枚举：草稿0、审批通过1、已作废2、审批中3。"],
        ],
        [3100, 6260],
        8.2,
    )
    heading(doc, "5.3 与前端的交互", 2)
    add_table(
        doc,
        ["页面/控件", "前端调用", "后端服务"],
        [
            ["登录页", "login(payload)", "POST /api/auth/login"],
            ["路由守卫", "检查 Pinia token", "无后端调用，未登录跳转 /login。"],
            ["列表查询区", "getReimbursements(query)", "GET /api/reimbursements"],
            ["基础数据下拉", "masterStore.load()", "GET /api/master/*"],
            ["报销表单初始化", "getReimbursement(id)", "GET /api/reimbursements/{id}"],
            ["保存草稿", "createDraft/saveDraft", "POST /drafts 或 PUT /{id}/draft"],
            ["提交", "persistDraft 后 submitReimbursement", "POST /{id}/submit"],
            ["审批/撤回/作废/复制/删除", "对应 api 函数", "approve/withdraw/void/copy/delete"],
        ],
        [2300, 3100, 3960],
        8.1,
    )
    heading(doc, "5.4 与第三方的交互", 2)
    paragraph(doc, "当前实现未接入第三方审批流、发票查验、票据识别、预算冻结或消息队列。后续如接入 BPM，可在 submit 后将状态保持为审批中，并通过适配器回写审批通过、驳回或作废状态。")


def add_key_points(doc: Document) -> None:
    heading(doc, "6 关键技术点", 1)
    add_table(
        doc,
        ["技术点", "实现方式", "注意事项"],
        [
            ["JWT鉴权", "HS256签名，payload含用户ID、员工ID、用户名、角色、签发/过期时间。", "生产环境需使用强随机 secret，通过环境变量配置。"],
            ["事务控制", "保存、提交、撤回、审批、作废、复制、删除使用 @Transactional。", "聚合保存采用替换子表策略，保证快照一致。"],
            ["补助计算", "按到达城市等级决定餐费标准，交通/通讯固定40元。", "实际金额不允许负数或超过标准。"],
            ["重复行程校验", "按同一出行人日期区间交叉判断。", "前后端双重校验，后端为最终准入规则。"],
            ["金额精度", "后端 BigDecimal 保留2位，比例保留6位。", "前端 roundMoney 保持展示与提交一致。"],
            ["权限控制", "列表非管理员过滤 owner_user_id；详情和写操作校验本人。", "后续审批角色可扩展单独权限规则。"],
            ["异常返回", "BusinessException 与全局异常处理统一 Result.failure。", "前端拦截器统一提示并处理401。"],
        ],
        [1900, 4300, 3160],
        8.0,
    )


def add_database_design(doc: Document, db: DbInfo) -> None:
    heading(doc, "7 数据库设计", 1)
    paragraph(doc, "数据库名：travel；连接端口：3306；本文档仅读取 information_schema 与各表 COUNT(*)，未执行任何建表、更新、删除或插入语句。")
    heading(doc, "7.1 表清单", 2)
    overview = [
        [table, TABLE_DESCRIPTIONS.get(table, ""), db.row_counts.get(table, 0)]
        for table in db.tables
    ]
    add_table(doc, ["表名", "中文说明", "当前行数"], overview, [3100, 4560, 1700], 8.2)
    heading(doc, "7.2 核心关系", 2)
    add_table(
        doc,
        ["关系", "说明"],
        [
            ["sys_company/sys_department -> sys_employee -> sys_user", "员工归属公司和部门，用户绑定员工。"],
            ["biz_business_type 自关联", "通过 parent_id 构成树形业务类型。"],
            ["fk_reim_main -> fk_reim_trip", "一张报销单可包含多条补录行程。"],
            ["fk_reim_trip -> fk_reim_subsidy -> fk_reim_subsidy_day", "一条行程对应一条补助汇总，补助汇总下有多日补助明细。"],
            ["fk_reim_main -> fk_reim_allocation", "一张报销单可包含多条费用归属分摊。"],
        ],
        [3000, 6360],
        8.4,
    )
    heading(doc, "7.3 表定义", 2)
    for table in db.tables:
        heading(doc, f"7.3 {table} - {TABLE_DESCRIPTIONS.get(table, '')}", 3)
        rows = []
        for col in db.columns.get(table, []):
            default = col.get("COLUMN_DEFAULT") or ""
            key = col.get("COLUMN_KEY") or ""
            extra = col.get("EXTRA") or ""
            rows.append([
                col["COLUMN_NAME"],
                col["COLUMN_TYPE"],
                "N" if col["IS_NULLABLE"] == "YES" else "Y",
                default,
                " ".join(v for v in [key, extra] if v),
                FIELD_DESCRIPTIONS.get(col["COLUMN_NAME"], ""),
            ])
        add_table(doc, ["字段名", "类型", "必填", "默认值", "约束", "中文说明"], rows, [2000, 1900, 850, 1450, 1450, 1710], 7.4)
    heading(doc, "7.4 索引设计", 2)
    idx_rows = []
    for table in db.tables:
        for idx in db.indexes.get(table, []):
            idx_rows.append([table, idx["INDEX_NAME"], "唯一" if idx["NON_UNIQUE"] == "0" else "普通", idx["columns"]])
    add_table(doc, ["表名", "索引名", "类型", "字段"], idx_rows, [2400, 3000, 1300, 2660], 7.9)
    heading(doc, "7.5 外键约束", 2)
    fk_rows = []
    for table in db.tables:
        for fk in db.fks.get(table, []):
            fk_rows.append([
                table,
                fk["COLUMN_NAME"],
                fk["CONSTRAINT_NAME"],
                f'{fk["REFERENCED_TABLE_NAME"]}.{fk["REFERENCED_COLUMN_NAME"]}',
            ])
    add_table(doc, ["表名", "字段", "约束名", "引用"], fk_rows, [2300, 1900, 3300, 1860], 8.0)


def add_api_docs(doc: Document) -> None:
    heading(doc, "8 接口文档", 1)
    paragraph(doc, "接口统一前缀为 /api。除登录接口外，请求需携带 Authorization: Bearer {token}。统一返回 Result<T>：code=0 表示成功，message=success，data 为业务数据；分页返回 PageResult<T>。")
    heading(doc, "8.1 接口总览", 2)
    add_table(
        doc,
        ["模块", "方法", "路径", "说明"],
        [
            ["认证", "POST", "/api/auth/login", "用户登录，返回JWT。"],
            ["基础数据", "GET", "/api/master/companies", "公司下拉。"],
            ["基础数据", "GET", "/api/master/departments", "部门下拉。"],
            ["基础数据", "GET", "/api/master/employees", "员工下拉。"],
            ["基础数据", "GET", "/api/master/business-types/tree", "业务类型树。"],
            ["基础数据", "GET", "/api/master/cities", "城市下拉与城市等级。"],
            ["基础数据", "GET", "/api/master/projects", "项目下拉。"],
            ["报销", "GET", "/api/reimbursements", "分页查询报销单列表。"],
            ["报销", "GET", "/api/reimbursements/{id}", "查询报销单详情。"],
            ["报销", "POST", "/api/reimbursements/drafts", "创建草稿。"],
            ["报销", "PUT", "/api/reimbursements/{id}/draft", "保存草稿。"],
            ["报销", "POST", "/api/reimbursements/{id}/submit", "提交报销单。"],
            ["报销", "POST", "/api/reimbursements/{id}/withdraw", "撤回报销单。"],
            ["报销", "POST", "/api/reimbursements/{id}/approve", "审批通过。"],
            ["报销", "POST", "/api/reimbursements/{id}/void", "作废报销单。"],
            ["报销", "POST", "/api/reimbursements/{id}/copy", "复制为新草稿。"],
            ["报销", "DELETE", "/api/reimbursements/{id}", "删除草稿。"],
        ],
        [1100, 900, 3600, 3760],
        7.8,
    )
    heading(doc, "8.2 公共响应", 2)
    add_table(
        doc,
        ["字段", "类型", "说明"],
        [
            ["code", "Integer", "0成功；400业务/参数失败；401认证失败；403权限不足；404资源不存在；500系统异常。"],
            ["message", "String", "成功为success，失败为可提示文本。"],
            ["data", "Object", "业务数据；无返回值动作可为空。"],
        ],
        [1900, 1700, 5760],
        8.4,
    )
    add_table(
        doc,
        ["分页字段", "类型", "说明"],
        [["total", "Long", "总条数"], ["current", "Long", "当前页"], ["size", "Long", "每页大小"], ["records", "List<T>", "当前页记录"]],
        [1900, 1700, 5760],
        8.4,
    )
    heading(doc, "8.3 登录接口", 2)
    add_table(
        doc,
        ["入参字段", "类型", "必传", "说明"],
        [["username", "String", "Y", "登录账号"], ["password", "String", "Y", "登录密码"]],
        [1900, 1600, 900, 4960],
        8.4,
    )
    add_table(
        doc,
        ["出参字段", "类型", "说明"],
        [["token", "String", "JWT令牌"], ["tokenType", "String", "固定Bearer"], ["expiresInMinutes", "Long", "有效期分钟，默认480"]],
        [1900, 1600, 5860],
        8.4,
    )
    heading(doc, "8.4 报销列表查询", 2)
    add_table(
        doc,
        ["参数", "类型", "必传", "说明"],
        [
            ["current", "Long", "N", "当前页，默认1。"],
            ["size", "Long", "N", "每页大小，默认10，最大100。"],
            ["reimNo", "String", "N", "报销单号模糊查询。"],
            ["title", "String", "N", "标题模糊查询。"],
            ["reason", "String", "N", "事由模糊查询。"],
            ["companyId", "Long", "N", "费用归属公司。"],
            ["departmentId", "Long", "N", "报销部门。"],
            ["reimburserId", "Long", "N", "报销人。"],
            ["businessTypeId", "Long", "N", "业务类型。"],
        ],
        [1900, 1400, 900, 5160],
        8.2,
    )
    add_table(
        doc,
        ["records字段", "类型", "说明"],
        [
            ["id/reimNo", "Long/String", "主键与报销单号。"],
            ["status/statusName", "Integer/String", "状态码与状态名称。"],
            ["reimburserDisplay/departmentDisplay", "String", "报销人与部门展示名。"],
            ["reimCompanyName/businessTypeName", "String", "费用归属公司与业务类型。"],
            ["reimbursementTitle/businessTripReason", "String", "标题与事由。"],
            ["subsidyTotal/creationTime", "Decimal/DateTime", "补助金额与创建时间。"],
        ],
        [2600, 1900, 4860],
        8.2,
    )
    heading(doc, "8.5 草稿保存入参", 2)
    add_table(
        doc,
        ["字段", "类型", "必传", "说明"],
        [
            ["reimbursementTitle", "String", "提交时Y", "报销标题，最长500字。"],
            ["reimburserId", "Long", "提交时Y", "报销人ID。"],
            ["reimDepartmentId", "Long", "提交时Y", "报销部门ID。"],
            ["reimCompanyId", "Long", "提交时Y", "费用归属公司ID。"],
            ["businessTypeId", "Long", "提交时Y", "业务类型ID。"],
            ["businessTripReason", "String", "提交时Y", "出差事由，最长500字。"],
            ["remarks", "String", "N", "备注，最长1000字。"],
            ["trips", "List<TripSaveDTO>", "提交时Y", "行程集合。"],
            ["trips[].travelerId", "Long", "Y", "出行人ID。"],
            ["trips[].departCityId/arriveCityId", "Long", "Y", "出发/到达城市ID。"],
            ["trips[].departDate/arriveDate", "LocalDate", "Y", "出发/到达日期。"],
            ["trips[].tripDescription", "String", "Y", "行程说明，最长500字。"],
            ["trips[].subsidyDays", "List<SubsidyDaySaveDTO>", "N", "补助日历选择与金额。"],
            ["allocations", "List<AllocationSaveDTO>", "提交时Y", "费用分摊集合。"],
            ["allocations[].companyId/projectId", "Long", "Y/N", "费用归属公司与项目。"],
            ["allocations[].allocationRatio", "Decimal", "Y", "分摊比例，0到1。"],
            ["allocations[].allocationAmount", "Decimal", "Y", "分摊金额。"],
        ],
        [2700, 1900, 1100, 3660],
        7.4,
    )
    heading(doc, "8.6 状态动作接口", 2)
    add_table(
        doc,
        ["接口", "入参", "出参", "说明"],
        [
            ["submit", "Path id", "ReimbursementActionVO", "校验通过后置为审批中。"],
            ["withdraw", "Path id", "ReimbursementActionVO", "审批中撤回为草稿。"],
            ["approve", "Path id", "ReimbursementActionVO", "审批中置为审批通过。"],
            ["void", "Path id", "ReimbursementActionVO", "未作废单据置为已作废。"],
            ["copy", "Path id", "ReimbursementDetailVO", "复制可见单据为新草稿。"],
            ["delete", "Path id", "Void", "删除本人草稿。"],
        ],
        [1900, 1600, 2400, 3460],
        8.2,
    )


WBS_ROWS = [
    ["1.0", "项目启动与范围确认", "确认模板、代码库、数据库、文档资料和交付物范围。", "项目范围说明、资料清单", "已完成"],
    ["1.1", "资料盘点", "读取模板、概要设计、API模板、表结构模板和WBS模板。", "资料差异清单", "已完成"],
    ["1.2", "现状核对", "核对后端、前端、数据库三端实现范围。", "现状分析记录", "已完成"],
    ["2.0", "需求分析", "梳理业务流程、状态、功能模块、页面控件和数据规则。", "需求拆解", "已完成"],
    ["2.1", "业务流程设计", "整理登录、查询、新增、保存、提交、审批、撤回、作废流程。", "业务流程图", "已完成"],
    ["2.2", "功能清单", "拆分认证、基础数据、列表、表单、行程、补助、分摊、状态操作。", "功能矩阵", "已完成"],
    ["3.0", "详细设计", "输出逻辑、框架图、时序图、数据库表定义和接口文档。", "详细设计报告", "已完成"],
    ["3.1", "接口设计", "整理 REST API、入参、出参、状态码和认证要求。", "接口文档", "已完成"],
    ["3.2", "数据库设计", "只读读取 travel 库表、字段、索引、外键、行数。", "表结构定义", "已完成"],
    ["4.0", "后端实现", "Spring Boot、Security、MyBatis-Plus、业务服务和测试。", "后端代码", "已完成"],
    ["4.1", "认证模块", "登录、JWT签发与过滤器。", "AuthController/AuthService", "已完成"],
    ["4.2", "基础数据模块", "公司、部门、员工、业务类型、城市、项目查询。", "MasterData模块", "已完成"],
    ["4.3", "报销核心模块", "列表、详情、草稿、提交、审批、撤回、作废、复制、删除。", "Reimbursement模块", "已完成"],
    ["5.0", "前端实现", "Vue页面、路由、Pinia、Axios、Element Plus控件。", "前端代码", "已完成"],
    ["5.1", "列表页", "查询表单、结果表格、分页和行操作。", "ReimbursementListView", "已完成"],
    ["5.2", "表单页", "基础信息、行程弹窗、补助日历、分摊和备注。", "ReimbursementFormView", "已完成"],
    ["6.0", "测试与联调", "单元测试、接口联调、前端构建和冒烟验证。", "测试记录", "进行中"],
    ["6.1", "后端测试", "登录、草稿、提交、审批、撤回、作废、重复行程测试。", "JUnit测试", "已完成"],
    ["6.2", "前端构建", "类型检查、打包和页面冒烟。", "构建结果", "待执行"],
    ["7.0", "交付与归档", "提交详细设计报告、WBS报告和最终交付说明。", "交付包", "进行中"],
]


def add_wbs_section(doc: Document, compact: bool = False) -> None:
    heading(doc, "9 WBS任务计划" if not compact else "3 WBS任务分解", 1)
    tasks = json.loads(WBS_TASKS_FILE.read_text(encoding="utf-8"))
    rows = [
        [
            task["code"],
            task["phase"],
            task["block"],
            task["taskName"],
            task["taskDesc"],
            task["output"],
            task["status"],
        ]
        for task in tasks
    ]
    add_table(
        doc,
        ["WBS编码", "计划阶段", "板块", "任务名称", "任务说明", "交付物", "状态"],
        rows,
        [850, 1250, 1100, 1800, 2700, 1100, 560],
        7.0,
    )


def add_risks(doc: Document, section_no: str = "10") -> None:
    heading(doc, f"{section_no} 风险与问题跟踪", 1)
    add_table(
        doc,
        ["风险/问题", "影响", "建议措施"],
        [
            ["审批接口已有实现，但未接入真实审批流或审批角色表。", "生产审批权限和流程节点可能不满足企业制度。", "后续引入审批角色、流程实例、审批意见和操作日志。"],
            ["草稿保存采用替换子表策略。", "并发编辑时后提交的数据可能覆盖先提交数据。", "增加 version 字段或更新时间并发校验。"],
            ["JWT secret 配置存在默认值。", "生产环境若使用默认值存在安全风险。", "部署时强制环境变量覆盖并定期轮换。"],
            ["基础数据暂无维护页面。", "字典数据变更依赖数据库脚本。", "补充后台维护或导入机制。"],
            ["接口文档为源码与数据库现状提取。", "若后续改动代码，需要同步更新文档。", "在交付前进行接口契约复核。"],
        ],
        [3000, 3000, 3360],
        8.1,
        WARN_FILL,
    )


def add_delivery_checklist(doc: Document) -> None:
    heading(doc, "附录 交付与自测清单", 1)
    add_table(
        doc,
        ["检查项", "结果", "说明"],
        [
            ["后端源码读取", "通过", "已读取 controller/service/dto/vo/security/db 脚本。"],
            ["前端源码读取", "通过", "已读取 router/api/store/list/form/types/utils。"],
            ["数据库只读核对", "通过", "仅读取 information_schema、COUNT(*)、状态统计。"],
            ["模板复用", "通过", "以差旅报销系统_详细设计WBS.docx 为文档底稿。"],
            ["接口覆盖", "通过", "覆盖认证、基础数据、报销列表/详情/草稿/状态动作。"],
            ["表结构覆盖", "通过", "覆盖 travel 库12张表、字段、索引、外键。"],
        ],
        [2900, 1100, 5360],
        8.3,
    )


def build_detail_report(db: DbInfo, diagrams: dict[str, Path]) -> Path:
    doc = Document(str(TEMPLATE))
    clear_body(doc)
    configure_document(doc, "差旅报销系统 | 详细设计报告")
    cover(
        doc,
        "详细设计报告",
        "逻辑设计、流程图、时序图、框架图、数据库表定义、接口文档与WBS",
        [
            ("项目", "费控云差旅报销模块"),
            ("版本", "V1.0"),
            ("日期", "2026-05-25"),
            ("后端路径", str(WORKSPACE)),
            ("前端路径", str(FRONTEND)),
            ("数据库", "MySQL 3306 / travel（只读核对）"),
            ("实现范围", "登录、基础数据、报销列表、报销单表单、补录行程、补助日历、费用分摊、草稿、提交、审批、撤回、作废、复制、删除草稿。"),
        ],
    )
    revision_table(doc)
    toc(
        doc,
        [
            "1 非功能性要求",
            "2 术语定义",
            "3 功能性需求描述",
            "4 功能详细设计",
            "5 技术实现设计",
            "6 关键技术点",
            "7 数据库设计",
            "8 接口文档",
            "9 WBS任务计划",
            "10 风险与问题跟踪",
            "附录 交付与自测清单",
        ],
    )
    page_break(doc)
    add_nonfunctional(doc)
    add_terms(doc)
    add_requirements(doc, diagrams, db)
    add_function_design(doc, diagrams)
    add_technical_design(doc, diagrams)
    add_key_points(doc)
    add_database_design(doc, db)
    add_api_docs(doc)
    add_wbs_section(doc)
    add_risks(doc)
    add_delivery_checklist(doc)
    doc.save(DETAIL_DOCX)
    return DETAIL_DOCX


def build_wbs_report(diagrams: dict[str, Path]) -> Path:
    doc = Document(str(TEMPLATE))
    clear_body(doc)
    configure_document(doc, "差旅报销系统 | WBS报告")
    cover(
        doc,
        "WBS报告",
        "工作分解结构、里程碑、交付物、责任建议、验收标准与风险跟踪",
        [
            ("项目", "费控云差旅报销模块"),
            ("版本", "V1.0"),
            ("日期", "2026-05-25"),
            ("编制依据", "详细设计模板、概要设计、前后端源码和数据库只读核对结果"),
            ("交付目标", "形成可执行的差旅报销系统任务拆解与验收计划。"),
        ],
    )
    revision_table(doc)
    toc(
        doc,
        [
            "1 WBS范围说明",
            "2 WBS框架图",
            "3 WBS任务分解",
            "4 里程碑与进度建议",
            "5 角色责任建议",
            "6 验收标准",
            "7 风险与问题跟踪",
        ],
    )
    page_break(doc)
    heading(doc, "1 WBS范围说明", 1)
    paragraph(doc, "本WBS覆盖差旅报销系统从资料盘点、需求分析、详细设计、后端实现、前端实现、数据库核对、测试联调到交付归档的完整工作拆解。当前代码已实现主要功能，本报告按“已完成/进行中/待执行”标记现状，并给出后续验收建议。")
    add_table(
        doc,
        ["范围项", "说明"],
        [
            ["包含", "登录鉴权、基础数据、报销列表、报销表单、行程、补助、分摊、状态流转、数据库表、接口文档、设计报告。"],
            ["不包含", "真实BPM审批流、票据识别、发票验真、预算冻结、支付结算、财务凭证、基础数据维护后台。"],
            ["交付物", "详细设计报告、WBS报告、源码、数据库脚本、接口说明、测试记录。"],
            ["约束", "数据库仅允许读取核对，不执行任何变更。"],
        ],
        [2200, 7160],
        8.5,
    )
    heading(doc, "2 WBS框架图", 1)
    add_image(doc, diagrams["wbs"], "图2-1 差旅报销系统WBS框架图")
    add_wbs_section(doc, compact=True)
    heading(doc, "4 里程碑与进度建议", 1)
    add_table(
        doc,
        ["里程碑", "建议日期", "准入/完成标准"],
        [
            ["M1 范围确认完成", "2026-05-25", "模板、资料、代码路径、数据库连接和交付物范围确认。"],
            ["M2 详细设计完成", "2026-05-25", "详细设计报告包含流程图、时序图、框架图、数据库表定义、接口文档和WBS。"],
            ["M3 后端验收", "2026-05-26", "后端单元测试通过，核心接口联调通过。"],
            ["M4 前端验收", "2026-05-27", "前端构建通过，登录、列表、表单、提交、状态操作冒烟通过。"],
            ["M5 交付归档", "2026-05-28", "报告、脚本、测试记录和部署说明归档。"],
        ],
        [2200, 1800, 5360],
        8.3,
    )
    heading(doc, "5 角色责任建议", 1)
    add_table(
        doc,
        ["角色", "主要责任", "关键交付"],
        [
            ["产品/业务", "确认差旅报销业务范围、状态、页面字段、审批规则。", "需求确认记录、验收用例。"],
            ["后端开发", "实现认证、基础数据、报销聚合服务、接口和数据库脚本。", "后端代码、接口说明、单元测试。"],
            ["前端开发", "实现登录、列表、表单、行程弹窗、补助日历、分摊交互。", "前端页面、构建产物。"],
            ["测试", "准备用例，执行接口、页面、状态流转、边界校验。", "测试报告、缺陷清单。"],
            ["运维/DBA", "准备运行环境、配置数据库连接和生产密钥。", "部署配置、数据库初始化记录。"],
        ],
        [1800, 4300, 3260],
        8.3,
    )
    heading(doc, "6 验收标准", 1)
    add_table(
        doc,
        ["验收域", "标准"],
        [
            ["登录与权限", "登录成功返回JWT；未登录不能访问业务接口；非本人不能查看或操作他人单据。"],
            ["列表查询", "按单号、标题、事由、公司、部门、报销人、业务类型分页查询正确。"],
            ["草稿保存", "主表、行程、补助日历、分摊可保存并回显；关闭未提交可保存草稿。"],
            ["提交校验", "必填、行程、日期重叠、分摊比例、分摊金额均按规则校验。"],
            ["状态流转", "草稿提交到审批中；审批中可撤回、通过、作废；草稿可删除；可复制为新草稿。"],
            ["数据库", "12张表结构、索引和外键与接口保存逻辑一致。"],
            ["文档", "详细设计报告和WBS报告内容完整，可指导开发、测试和交付。"],
        ],
        [2200, 7160],
        8.3,
    )
    add_risks(doc, "7")
    doc.save(WBS_DOCX)
    return WBS_DOCX


def main() -> None:
    diagrams = {
        "architecture": save_architecture_diagram(),
        "business": save_business_flow_diagram(),
        "sequence": save_sequence_diagram(),
        "wbs": save_wbs_diagram(),
    }
    db = load_db_info()
    detail = build_detail_report(db, diagrams)
    wbs = build_wbs_report(diagrams)
    print(detail)
    print(wbs)


if __name__ == "__main__":
    main()
