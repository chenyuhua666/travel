import os
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TARGET = Path(os.environ["TARGET_PATH"])

TOC_ENTRIES = [
    (0, "目录", "1"),
    (0, "1 非功能性要求", "3"),
    (0, "2 术语定义", "3"),
    (0, "3 功能性需求描述", "3"),
    (1, "3.1 需求概述", "3"),
    (1, "3.2 用户角色", "4"),
    (1, "3.3 单据状态", "4"),
    (1, "3.4 业务流程", "4"),
    (1, "3.5 功能范围", "5"),
    (0, "4 功能详细设计", "5"),
    (1, "4.1 登录与鉴权", "5"),
    (1, "4.2 基础数据", "6"),
    (1, "4.3 报销单列表", "6"),
    (1, "4.4 报销单表单与基础信息", "6"),
    (1, "4.5 补录行程", "7"),
    (1, "4.6 补助信息与补助日历", "7"),
    (1, "4.7 费用合计与费用分摊", "7"),
    (1, "4.8 保存草稿与提交", "8"),
    (1, "4.9 状态操作", "8"),
    (1, "4.10 异常处置", "9"),
    (0, "5 技术实现设计", "9"),
    (1, "5.1 系统结构设计", "9"),
    (1, "5.2 接口及核心类设计", "10"),
    (1, "5.3 与前端的交互", "10"),
    (1, "5.4 与第三方的交互", "10"),
    (0, "6 关键技术点", "11"),
    (0, "7 数据库设计", "11"),
    (1, "7.1 表清单", "11"),
    (1, "7.2 核心关系", "12"),
    (1, "7.3 表定义", "12"),
    (1, "7.4 索引设计", "18"),
    (1, "7.5 外键约束", "19"),
    (0, "8 接口文档", "19"),
    (1, "8.1 接口总览", "19"),
    (1, "8.2 公共响应", "20"),
    (1, "8.3 登录接口", "20"),
    (1, "8.4 报销列表查询", "21"),
    (1, "8.5 草稿保存入参", "21"),
    (1, "8.6 状态动作接口", "22"),
    (0, "9 WBS任务计划", "22"),
    (0, "10 风险与问题跟踪", "23"),
    (0, "附录 交付与自测清单", "24"),
]


def set_run_font(run, size=10.5, name="宋体", color=None, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def find_toc_heading(doc):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == "目录":
            return paragraph
    raise RuntimeError("TOC heading not found")


def find_first_body_heading(doc):
    seen_toc = False
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == "目录":
            seen_toc = True
            continue
        if seen_toc and paragraph.text.strip() == "1 非功能性要求":
            return paragraph
    raise RuntimeError("first body heading not found")


def remove_between(start_element, end_element):
    current = start_element.getnext()
    while current is not None and current is not end_element:
        nxt = current.getnext()
        current.getparent().remove(current)
        current = nxt


def make_toc_paragraph(doc, level, title, page):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.left_indent = Pt(22 * level)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.55), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    run = paragraph.add_run(f"{title}\t{page}")
    set_run_font(run, size=10.5)
    return paragraph


def main():
    doc = Document(str(TARGET))
    heading = find_toc_heading(doc)
    first_body = find_first_body_heading(doc)
    remove_between(heading._element, first_body._element)
    new_elements = [make_toc_paragraph(doc, *entry)._element for entry in TOC_ENTRIES]
    page_break = doc.add_paragraph()
    page_break.add_run().add_break(WD_BREAK.PAGE)
    new_elements.append(page_break._element)
    for element in new_elements:
        first_body._element.addprevious(element)
    doc.save(str(TARGET))
    print(str(TARGET))


if __name__ == "__main__":
    main()
