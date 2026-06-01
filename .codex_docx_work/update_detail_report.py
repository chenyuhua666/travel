import os
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"E:\codelocation\IDEA\travel")
WORK = ROOT / ".codex_docx_work"
TARGET = Path(os.environ.get("TARGET_PATH", r"E:\qq下载路径\差旅报销系统_详细设计报告_更新版.docx"))
TEMPLATE = Path(r"E:\qq下载路径\项目文档资料\详细设计规范模板.docx")
LOGO = WORK / "media" / "DOC3" / "image1.png"


WBS_MODULES = [
    ("差旅报销系统", "后端服务", "登录认证接口", "新增 /api/auth/login，完成用户名密码校验、PBKDF2 密码摘要比对、JWT 签发与 Bearer Token 鉴权。"),
    ("差旅报销系统", "后端服务", "主数据接口", "新增公司、部门、员工、业务类型树、城市、项目查询接口，为前端下拉、级联和筛选提供基础数据。"),
    ("差旅报销系统", "后端服务", "报销单聚合接口", "新增列表、详情、创建草稿、保存草稿、提交、撤回、审批通过、作废、复制、删除草稿等接口。"),
    ("差旅报销系统", "后端服务", "补助与分摊服务", "新增行程日期重叠校验、逐日补助计算、餐费/交通/通讯补助选择、分摊比例与金额一致性校验。"),
    ("差旅报销系统", "后端服务", "异常与安全处理", "统一 Result<T> 响应、GlobalExceptionHandler 异常封装、Spring Security 无状态认证和本人数据权限校验。"),
    ("差旅报销系统", "前端页面", "登录页", "新增登录表单、Token 本地保存、请求拦截器、401 自动清理并跳转登录页。"),
    ("差旅报销系统", "前端页面", "报销列表", "新增查询条件、分页表格、新增、编辑、审批、作废、复制、删除等入口和状态化按钮。"),
    ("差旅报销系统", "前端页面", "报销表单", "新增基础信息、补录行程弹窗、补助日历、费用合计、分摊区域、备注、保存草稿和提交动作。"),
    ("差旅报销系统", "前端页面", "状态操作", "新增草稿提交、审批中撤回、审批通过、作废、草稿删除、复制为新草稿的前后端交互。"),
    ("差旅报销系统", "数据库", "用户与基础数据", "新增 sys_company、sys_department、sys_employee、sys_user、biz_business_type、biz_city、biz_project 等基础表。"),
    ("差旅报销系统", "数据库", "报销业务数据", "新增 fk_reim_main、fk_reim_trip、fk_reim_subsidy、fk_reim_subsidy_day、fk_reim_allocation 业务表。"),
    ("差旅报销系统", "数据库", "约束与级联", "补充主键、唯一键、外键、索引与 ON DELETE CASCADE，保证主子表保存和删除一致。"),
    ("差旅报销系统", "测试联调", "后端测试", "覆盖登录、草稿保存、提交审批、撤回、作废、复制、删除、重复行程拦截等核心用例。"),
    ("差旅报销系统", "测试联调", "前端验证", "执行 Vue/Vite 构建，验证登录、列表查询、表单保存、提交、状态操作和页面回显。"),
    ("差旅报销系统", "测试联调", "接口联调", "核对 JWT、主数据、列表、详情、草稿保存、状态流转接口契约，确保前后端字段一致。"),
    ("差旅报销系统", "交付归档", "文档交付", "归档详细设计报告、接口说明、表结构定义、WBS 报告、测试记录和部署说明。"),
]


def ensure_logo():
    LOGO.parent.mkdir(parents=True, exist_ok=True)
    if LOGO.exists():
        return
    with zipfile.ZipFile(TEMPLATE) as zf:
        LOGO.write_bytes(zf.read("word/media/image1.png"))


def set_run_font(run, size=None, bold=None, color=None, name="宋体"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def paragraph_text(paragraph, text, size=10.5, bold=False, color=None, align=None, name="宋体"):
    paragraph.text = ""
    if align is not None:
        paragraph.alignment = align
    for idx, part in enumerate(str(text).split("\n")):
        if idx:
            paragraph.add_run().add_break()
        run = paragraph.add_run(part)
        set_run_font(run, size=size, bold=bold, color=color, name=name)
    return paragraph


def add_page_field(paragraph):
    paragraph.text = ""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(text)
    run._r.append(fld_end)
    set_run_font(run, size=9, name="宋体")


def add_paragraph_bottom_border(paragraph, color="9E9E9E", size="6"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)


def clear_part(part):
    for child in list(part._element):
        part._element.remove(child)
    return part.add_paragraph()


def set_header_footer(doc):
    for section in doc.sections:
        section.header_distance = Cm(0.7)
        section.footer_distance = Cm(0.7)
        h_para = clear_part(section.header)
        h_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        h_para.paragraph_format.space_after = Pt(0)
        h_run = h_para.add_run()
        h_run.add_picture(str(LOGO), width=Inches(0.82))
        add_paragraph_bottom_border(h_para)

        f_para = clear_part(section.footer)
        add_page_field(f_para)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="000000", sz="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), sz)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, size=9, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.LEFT, name="宋体"):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph_text(paragraph, text, size=size, bold=bold, color=color, align=align, name=name)


def set_table_grid(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.allow_autofit = False
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell)
            set_cell_margins(cell)


def set_col_widths(table, widths_cm):
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            cell = row.cells[idx]
            cell.width = Cm(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 567)))
            tc_w.set(qn("w:type"), "dxa")


def find_first_nonblank_paragraph(doc, text):
    for para in doc.paragraphs:
        if para.text.strip() == text:
            return para
    raise RuntimeError(f"paragraph not found: {text}")


def remove_elements_before(element):
    previous = element.getprevious()
    while previous is not None:
        before = previous.getprevious()
        element.getparent().remove(previous)
        previous = before


def move_elements_before(elements, reference_element):
    for element in elements:
        reference_element.addprevious(element)


def build_cover_elements(doc):
    elements = []

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(24)
    title.paragraph_format.space_after = Pt(20)
    paragraph_text(title, "详细设计规范", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, name="黑体")
    elements.append(title._p)

    meta = doc.add_table(rows=5, cols=2)
    set_table_grid(meta)
    set_col_widths(meta, [7.8, 9.0])
    left = meta.cell(0, 0).merge(meta.cell(4, 0))
    set_cell_text(left, "文件状态：\n\n        草稿        □\n\n        修改        □\n\n        正式发布    ■", size=10, name="宋体")
    right_values = [
        "所属项目编号：TRAVEL-LOCAL-2026",
        "版    本：V1.0",
        "撰 写 人：项目开发组",
        "完成日期：2026-05-25",
        "发布日期：2026-05-25",
    ]
    for idx, value in enumerate(right_values):
        set_cell_text(meta.cell(idx, 1), value, size=10, name="宋体")
    elements.append(meta._tbl)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(12)
    elements.append(spacer._p)

    rev = doc.add_table(rows=2, cols=6)
    set_table_grid(rev)
    set_col_widths(rev, [1.5, 1.5, 1.8, 4.6, 3.1, 4.3])
    rev_rows = [
        ["序号", "类别", "版本", "作者", "时间", "备注"],
        ["1", "新增", "1.0", "项目开发组", "2026-05-25", "新增"],
    ]
    for r_idx, row in enumerate(rev_rows):
        for c_idx, value in enumerate(row):
            set_cell_text(
                rev.cell(r_idx, c_idx),
                value,
                size=9,
                bold=(r_idx == 0),
                align=WD_ALIGN_PARAGRAPH.CENTER if c_idx != 5 else WD_ALIGN_PARAGRAPH.LEFT,
                name="宋体",
            )
    elements.append(rev._tbl)

    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(34)
    elements.append(gap._p)
    return elements


def replace_cover(doc):
    toc_para = find_first_nonblank_paragraph(doc, "目录")
    ref = toc_para._p
    remove_elements_before(ref)
    cover_elements = build_cover_elements(doc)
    move_elements_before(cover_elements, ref)


def make_wbs_module_table(doc):
    table = doc.add_table(rows=len(WBS_MODULES) + 1, cols=4)
    set_table_grid(table)
    set_col_widths(table, [2.0, 2.4, 3.1, 9.3])
    headers = ["系统", "一级模块", "功能点", "修改功能点"]
    for idx, value in enumerate(headers):
        cell = table.cell(0, idx)
        set_cell_shading(cell, "FFF200")
        set_cell_text(cell, value, size=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, name="宋体")

    for row_idx, (_, module, feature, change) in enumerate(WBS_MODULES, start=1):
        values = ["", module, feature, change]
        for col_idx, value in enumerate(values):
            cell = table.cell(row_idx, col_idx)
            set_cell_shading(cell, "E2F0D9")
            set_cell_text(
                cell,
                value,
                size=7.5,
                color="D60000" if col_idx == 1 else None,
                bold=(col_idx == 1),
                align=WD_ALIGN_PARAGRAPH.CENTER if col_idx < 3 else WD_ALIGN_PARAGRAPH.LEFT,
                name="宋体",
            )

    system_cell = table.cell(1, 0).merge(table.cell(len(WBS_MODULES), 0))
    set_cell_shading(system_cell, "E2F0D9")
    set_cell_text(system_cell, "差旅报销\n系统", size=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, name="宋体")

    start = 1
    while start <= len(WBS_MODULES):
        module = WBS_MODULES[start - 1][1]
        end = start
        while end <= len(WBS_MODULES) and WBS_MODULES[end - 1][1] == module:
            end += 1
        merged = table.cell(start, 1).merge(table.cell(end - 1, 1))
        set_cell_shading(merged, "E2F0D9")
        set_cell_text(merged, module, size=8, bold=True, color="D60000", align=WD_ALIGN_PARAGRAPH.CENTER, name="宋体")
        start = end

    return table


def replace_wbs_table(doc):
    old = None
    for table in doc.tables:
        first = [cell.text.strip() for cell in table.rows[0].cells]
        if first[:4] == ["WBS编码", "计划阶段", "板块", "任务名称"]:
            old = table
            break
    if old is None:
        raise RuntimeError("old WBS table not found")
    new = make_wbs_module_table(doc)
    old._tbl.addprevious(new._tbl)
    old._tbl.getparent().remove(old._tbl)


def main():
    ensure_logo()
    doc = Document(str(TARGET))
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    set_header_footer(doc)
    replace_cover(doc)
    replace_wbs_table(doc)

    doc.core_properties.title = "差旅报销系统详细设计报告"
    doc.core_properties.subject = "详细设计报告"
    doc.core_properties.author = "项目开发组"
    doc.save(str(TARGET))
    print(str(TARGET))


if __name__ == "__main__":
    main()
