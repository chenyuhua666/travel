import json
import os
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"E:\codelocation\IDEA\travel")
WORK = ROOT / ".codex_docx_work"
TARGET = Path(r"E:\qq下载路径\差旅报销系统_WBS报告_更新版.docx")
TEMPLATE = Path(r"E:\qq下载路径\项目文档资料\详细设计规范模板.docx")
LOGO = WORK / "media" / "DOC3" / "image1.png"
WBS_IMAGE = WORK / "media" / "TARGET" / "image1.png"
TOC_JSON = WORK / "toc_pages.json"


SCOPE_FALLBACK = [
    ["范围项", "说明"],
    ["包含", "登录鉴权、基础数据、报销列表、报销表单、行程、补助、分摊、状态流转、数据库表、接口文档、设计报告。"],
    ["不包含", "真实BPM审批流、票据识别、发票验真、预算冻结、支付结算、财务凭证、基础数据维护后台。"],
    ["交付物", "详细设计报告、WBS报告、源码、数据库脚本、接口说明、测试记录。"],
    ["约束", "数据库仅允许读取核对，不执行任何变更。"],
]

MILESTONE_FALLBACK = [
    ["里程碑", "建议日期", "准入/完成标准"],
    ["M1 范围确认完成", "2026-05-25", "模板、资料、代码路径、数据库连接和交付物范围确认。"],
    ["M2 详细设计完成", "2026-05-25", "详细设计报告包含流程图、时序图、框架图、数据库表定义、接口文档和WBS。"],
    ["M3 后端验收", "2026-05-26", "后端单元测试通过，核心接口联调通过。"],
    ["M4 前端验收", "2026-05-27", "前端构建通过，登录、列表、表单、提交、状态操作冒烟通过。"],
    ["M5 交付归档", "2026-05-28", "报告、脚本、测试记录和部署说明归档。"],
]

ROLE_FALLBACK = [
    ["角色", "主要责任", "关键交付"],
    ["产品/业务", "确认差旅报销业务范围、状态、页面字段、审批规则。", "需求确认记录、验收用例。"],
    ["后端开发", "实现认证、基础数据、报销聚合服务、接口和数据库脚本。", "后端代码、接口说明、单元测试。"],
    ["前端开发", "实现登录、列表、表单、行程弹窗、补助日历、分摊交互。", "前端页面、构建产物。"],
    ["测试", "准备用例，执行接口、页面、状态流转、边界校验。", "测试报告、缺陷清单。"],
    ["运维/DBA", "准备运行环境、配置数据库连接和生产密钥。", "部署配置、数据库初始化记录。"],
]

ACCEPTANCE_FALLBACK = [
    ["验收域", "标准"],
    ["登录与权限", "登录成功返回JWT；未登录不能访问业务接口；非本人不能查看或操作他人单据。"],
    ["列表查询", "按单号、标题、事由、公司、部门、报销人、业务类型分页查询正确。"],
    ["草稿保存", "主表、行程、补助日历、分摊可保存并回显；关闭未提交可保存草稿。"],
    ["提交校验", "必填、行程、日期重叠、分摊比例、分摊金额均按规则校验。"],
    ["状态流转", "草稿提交到审批中；审批中可撤回、通过、作废；草稿可删除；可复制为新草稿。"],
    ["数据库", "12张表结构、索引和外键与接口保存逻辑一致。"],
    ["文档", "详细设计报告和WBS报告内容完整，可指导开发、测试和交付。"],
]

RISK_FALLBACK = [
    ["风险/问题", "影响", "建议措施"],
    ["审批接口已有实现，但未接入真实审批流或审批角色表。", "生产审批权限和流程节点可能不满足企业制度。", "后续引入审批角色、流程实例、审批意见和操作日志。"],
    ["草稿保存采用替换子表策略。", "并发编辑时后提交的数据可能覆盖先提交数据。", "增加 version 字段或更新时间并发校验。"],
    ["JWT secret 配置存在默认值。", "生产环境若使用默认值存在安全风险。", "部署时强制环境变量覆盖并定期轮换。"],
    ["基础数据暂无维护页面。", "字典数据变更依赖数据库脚本。", "补充后台维护或导入机制。"],
    ["接口文档为源码与数据库现状提取。", "若后续改动代码，需要同步更新文档。", "在交付前进行接口契约复核。"],
]

WBS_MODULES = [
    ("差旅报销系统", "后端服务", "登录认证接口", "新增 /api/auth/login，完成用户名密码校验、PBKDF2 密码摘要比对、JWT 签发与 Bearer Token 鉴权。"),
    ("差旅报销系统", "后端服务", "主数据接口", "新增公司、部门、员工、业务类型树、城市、项目查询接口，为前端下拉、级联和筛选提供基础数据。"),
    ("差旅报销系统", "后端服务", "报销单聚合接口", "新增列表、详情、创建草稿、保存草稿、提交、撤回、审批通过、作废、复制、删除草稿等接口。"),
    ("差旅报销系统", "后端服务", "补助与分摊服务", "新增行程日期重叠校验、逐日补助计算、餐费/交通/通讯补助选择、分摊比例与金额一致性校验。"),
    ("差旅报销系统", "后端服务", "异常与安全处理", "统一 Result<T> 响应、GlobalExceptionHandler 异常封装、Spring Security 无状态认证和本人数据权限校验。"),
    ("差旅报销系统", "前端页面", "登录页", "新增登录表单、Token 本地保存、请求拦截器、401 自动清理并跳转登录页。"),
    ("差旅报销系统", "前端页面", "报销列表", "新增查询条件、分页表格、新增、编辑、审批、作废、复制、删除等入口和状态化按钮。"),
    ("差旅报销系统", "前端页面", "报销表单", "新增基础信息、补录行程弹窗、补助日历、费用合计、分摊区域、备注、保存草稿和提交动作。"),
    ("差旅报销系统", "前端页面", "状态操作", "新增草稿提交、审批中撤回、审批通过、作废、草稿删除、复制为新草稿的前后端交互。"),
    ("差旅报销系统", "数据库", "用户与基础数据", "新增 sys_company、sys_department、sys_employee、sys_user、biz_business_type、biz_city、biz_project 等基础表。"),
    ("差旅报销系统", "数据库", "报销业务数据", "新增 fk_reim_main、fk_reim_trip、fk_reim_subsidy、fk_reim_subsidy_day、fk_reim_allocation 业务表。"),
    ("差旅报销系统", "数据库", "约束与级联", "补充主键、唯一键、外键、索引与 ON DELETE CASCADE，保证主子表保存和删除一致。"),
    ("差旅报销系统", "测试联调", "后端测试", "覆盖登录、草稿保存、提交审批、撤回、作废、复制、删除、重复行程拦截等核心用例。"),
    ("差旅报销系统", "测试联调", "前端验证", "执行 Vue/Vite 构建，验证登录、列表查询、表单保存、提交、状态操作和页面回显。"),
    ("差旅报销系统", "测试联调", "接口联调", "核对 JWT、主数据、列表、详情、草稿保存、状态流转接口契约，确保前后端字段一致。"),
    ("差旅报销系统", "交付归档", "文档交付", "归档详细设计报告、接口说明、表结构定义、WBS 报告、测试记录和部署说明。"),
]


def ensure_assets():
    WORK.mkdir(parents=True, exist_ok=True)
    LOGO.parent.mkdir(parents=True, exist_ok=True)
    if not LOGO.exists():
        with zipfile.ZipFile(TEMPLATE) as zf:
            LOGO.write_bytes(zf.read("word/media/image1.png"))
    if not WBS_IMAGE.exists() and TARGET.exists():
        WBS_IMAGE.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(TARGET) as zf:
            WBS_IMAGE.write_bytes(zf.read("word/media/image1.png"))


def set_run_font(run, size=None, bold=None, color=None, name="宋体", italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def paragraph_text(paragraph, text, size=10.5, bold=False, color=None, align=None, name="宋体"):
    paragraph.text = ""
    if align is not None:
        paragraph.alignment = align
    for idx, part in enumerate(str(text).split("\n")):
        if idx:
            paragraph.add_run().add_break()
        run = paragraph.add_run(part)
        set_run_font(run, size=size, bold=bold, color=color, name=name)
    return paragraph


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.text = ""
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(text)
    run._r.append(fld_end)
    set_run_font(run, size=9, name="宋体")


def add_paragraph_bottom_border(paragraph, color="9E9E9E", size="6"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="000000", sz="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), sz)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    values = {"top": top, "start": start, "bottom": bottom, "end": end}
    for name, value in values.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, size=9, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.LEFT, name="宋体"):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph_text(paragraph, text, size=size, bold=bold, color=color, align=align, name=name)


def set_table_grid(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.allow_autofit = False
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell)
            set_cell_margins(cell)


def set_col_widths(table, widths_cm):
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            cell = row.cells[idx]
            cell.width = Cm(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 567)))
            tc_w.set(qn("w:type"), "dxa")


def add_simple_table(doc, data, widths_cm, header_fill="E8F1DF", body_fill=None, font_size=9):
    table = doc.add_table(rows=len(data), cols=len(data[0]))
    set_table_grid(table)
    set_col_widths(table, widths_cm)
    for r_idx, row in enumerate(data):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            if r_idx == 0:
                set_cell_shading(cell, header_fill)
                set_cell_text(cell, value, size=font_size, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, name="宋体")
            else:
                if body_fill:
                    set_cell_shading(cell, body_fill)
                align = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 and len(row) <= 3 else WD_ALIGN_PARAGRAPH.LEFT
                set_cell_text(cell, value, size=font_size, align=align, name="宋体")
    doc.add_paragraph()
    return table


def table_to_matrix(table):
    matrix = []
    for row in table.rows:
        matrix.append([cell.text.replace("\n", " / ").strip() for cell in row.cells])
    return matrix


def find_table(doc, first_row):
    for table in doc.tables:
        values = [cell.text.strip() for cell in table.rows[0].cells]
        if values[: len(first_row)] == first_row:
            return table_to_matrix(table)
    return None


def load_existing_tables():
    if not TARGET.exists():
        return {}
    src = Document(str(TARGET))
    return {
        "scope": find_table(src, ["范围项", "说明"]) or SCOPE_FALLBACK,
        "milestone": find_table(src, ["里程碑", "建议日期", "准入/完成标准"]) or MILESTONE_FALLBACK,
        "role": find_table(src, ["角色", "主要责任", "关键交付"]) or ROLE_FALLBACK,
        "acceptance": find_table(src, ["验收域", "标准"]) or ACCEPTANCE_FALLBACK,
        "risk": find_table(src, ["风险/问题", "影响", "建议措施"]) or RISK_FALLBACK,
    }


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading("", level=level)
    paragraph.paragraph_format.space_before = Pt(8 if level == 1 else 5)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    set_run_font(run, size=14 if level == 1 else 12, bold=True, color="000000", name="黑体")
    return paragraph


def add_body_paragraph(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Pt(21)
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph_text(paragraph, text, size=10.5, name="宋体")
    return paragraph


def load_toc_entries():
    default = [
        ["1 WBS范围说明", 2],
        ["2 WBS框架图", 2],
        ["3 WBS模块清单", 3],
        ["4 里程碑与进度建议", 4],
        ["5 角色责任建议", 4],
        ["6 验收标准", 5],
        ["7 风险与问题跟踪", 5],
    ]
    if TOC_JSON.exists():
        data = json.loads(TOC_JSON.read_text(encoding="utf-8"))
        return [[item["title"], item["page"]] for item in data]
    return default


def add_cover(doc):
    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(28)
    title.paragraph_format.space_after = Pt(24)
    paragraph_text(title, "差旅报销系统 WBS报告", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, name="黑体")

    meta = doc.add_table(rows=5, cols=2)
    set_table_grid(meta)
    set_col_widths(meta, [8.2, 8.6])
    left = meta.cell(0, 0).merge(meta.cell(4, 0))
    set_cell_text(left, "文件状态：\n\n        草稿        □\n\n        修改        □\n\n        正式发布    ■", size=10, name="宋体")
    right_values = [
        "所属项目编号：TRAVEL-LOCAL-2026",
        "版    本：V1.0",
        "撰 写 人：项目开发组",
        "完成日期：2026-05-25",
        "发布日期：2026-05-25",
    ]
    for idx, value in enumerate(right_values):
        set_cell_text(meta.cell(idx, 1), value, size=10, name="宋体")

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(12)

    rev = doc.add_table(rows=2, cols=6)
    set_table_grid(rev)
    set_col_widths(rev, [1.5, 1.5, 1.8, 4.6, 3.1, 4.3])
    rev_rows = [
        ["序号", "类别", "版本", "作者", "时间", "备注"],
        ["1", "新增", "1.0", "项目开发组", "2026-05-25", "新增"],
    ]
    for r_idx, row in enumerate(rev_rows):
        for c_idx, value in enumerate(row):
            set_cell_text(
                rev.cell(r_idx, c_idx),
                value,
                size=9,
                bold=(r_idx == 0),
                align=WD_ALIGN_PARAGRAPH.CENTER if c_idx != 5 else WD_ALIGN_PARAGRAPH.LEFT,
                name="宋体",
            )

    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(70)

    toc_title = doc.add_paragraph()
    toc_title.paragraph_format.space_after = Pt(8)
    paragraph_text(toc_title, "目录", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, name="黑体")

    entries = [["目录", 1], *load_toc_entries()]
    for title, page in entries:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.65), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        run = p.add_run(f"{title}\t{page}")
        set_run_font(run, size=9.5, name="宋体")

    doc.add_page_break()


def add_wbs_module_table(doc):
    table = doc.add_table(rows=len(WBS_MODULES) + 1, cols=4)
    set_table_grid(table)
    set_col_widths(table, [2.0, 2.4, 3.1, 9.3])
    headers = ["系统", "一级模块", "功能点", "修改功能点"]
    for idx, value in enumerate(headers):
        cell = table.cell(0, idx)
        set_cell_shading(cell, "FFF200")
        set_cell_text(cell, value, size=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, name="宋体")

    for row_idx, (_, module, feature, change) in enumerate(WBS_MODULES, start=1):
        values = ["", module, feature, change]
        for col_idx, value in enumerate(values):
            cell = table.cell(row_idx, col_idx)
            set_cell_shading(cell, "E2F0D9")
            set_cell_text(
                cell,
                value,
                size=7.5,
                color="D60000" if col_idx == 1 else None,
                bold=(col_idx == 1),
                align=WD_ALIGN_PARAGRAPH.CENTER if col_idx < 3 else WD_ALIGN_PARAGRAPH.LEFT,
                name="宋体",
            )

    system_cell = table.cell(1, 0).merge(table.cell(len(WBS_MODULES), 0))
    set_cell_shading(system_cell, "E2F0D9")
    set_cell_text(system_cell, "差旅报销\n系统", size=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, name="宋体")

    start = 1
    while start <= len(WBS_MODULES):
        module = WBS_MODULES[start - 1][1]
        end = start
        while end <= len(WBS_MODULES) and WBS_MODULES[end - 1][1] == module:
            end += 1
        merged = table.cell(start, 1).merge(table.cell(end - 1, 1))
        set_cell_shading(merged, "E2F0D9")
        set_cell_text(merged, module, size=8, bold=True, color="D60000", align=WD_ALIGN_PARAGRAPH.CENTER, name="宋体")
        start = end

    doc.add_paragraph()
    return table


def build_document():
    ensure_assets()
    old_tables = load_existing_tables()

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)

    header = section.header
    h_para = header.paragraphs[0]
    h_para.paragraph_format.space_after = Pt(0)
    h_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h_run = h_para.add_run()
    h_run.add_picture(str(LOGO), width=Inches(0.82))
    add_paragraph_bottom_border(h_para)

    footer = section.footer
    add_page_field(footer.paragraphs[0])

    add_cover(doc)

    add_heading(doc, "1 WBS范围说明", level=1)
    add_body_paragraph(
        doc,
        "本WBS覆盖差旅报销系统从资料盘点、需求分析、详细设计、后端实现、前端实现、数据库核对、测试联调到交付归档的完整工作拆解。当前代码已实现主要功能，本报告按项目交付链路整理模块、任务、责任和验收建议。",
    )
    add_simple_table(doc, old_tables["scope"], [2.8, 14.0], font_size=9)

    add_heading(doc, "2 WBS框架图", level=1)
    if WBS_IMAGE.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(WBS_IMAGE), width=Cm(16.5))
        cap = doc.add_paragraph()
        paragraph_text(cap, "图2-1 差旅报销系统WBS框架图", size=9, align=WD_ALIGN_PARAGRAPH.CENTER, name="宋体")
    doc.add_page_break()

    add_heading(doc, "3 WBS模块清单", level=1)
    add_body_paragraph(
        doc,
        "本节按照“系统 / 一级模块 / 功能点 / 修改功能点”的样式组织WBS模块，便于开发、测试和验收按模块扫描工作范围。",
    )
    add_wbs_module_table(doc)
    doc.add_page_break()

    add_heading(doc, "4 里程碑与进度建议", level=1)
    add_simple_table(doc, old_tables["milestone"], [4.0, 3.0, 9.8], font_size=9)

    add_heading(doc, "5 角色责任建议", level=1)
    add_simple_table(doc, old_tables["role"], [3.1, 7.0, 6.7], font_size=9)

    add_heading(doc, "6 验收标准", level=1)
    add_simple_table(doc, old_tables["acceptance"], [3.3, 13.5], font_size=9)

    doc.add_page_break()
    add_heading(doc, "7 风险与问题跟踪", level=1)
    add_simple_table(doc, old_tables["risk"], [6.2, 5.2, 5.4], font_size=8.5)

    doc.core_properties.title = "差旅报销系统 WBS报告"
    doc.core_properties.subject = "WBS报告"
    doc.core_properties.author = "项目开发组"
    doc.save(str(TARGET))
    print(str(TARGET))


if __name__ == "__main__":
    build_document()
